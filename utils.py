# -*- coding: utf-8 -*-

import os
import re
import time
import math
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from state import AgentState, require_state_value
from openalex_client import get_openalex_author_metrics
from semantic_scholar_client import get_semantic_scholar_author_metrics
from expert_enrichment import enrich_expert_details
from safe_logging import safe_print
from source_registry import source_records_for_query
from survival_verification import verify_survival_status
from chinese_output import translate_delivery_dataframe

EXPERT_TABLE_HEADERS = [
    "专家姓名", "国籍", "个人主页", "邮箱/电话", "研究兴趣",
    "工作单位", "职位", "工作经历", "教育背景", "H指数",
    "主要成果", "国内合作学者与单位", "入选依据", "领域关联依据", "生存状态", "信息来源"
]

CORE_DELIVERY_COLUMNS = [
    "姓名",
    "国籍",
    "所在工作机构",
    "职位",
    "学科领域",
    "细分领域",
    "工作经历",
    "主要成果",
    "联系方式",
    "国内合作学者",
    "国内合作单位",
    "入选依据",
]

SUPPORTING_DELIVERY_COLUMNS = [
    "个人主页",
    "研究兴趣",
    "教育背景",
    "性别",
    "语言",
    "H指数",
    "G指数",
    "i10指数",
    "总被引次数",
    "评价_H指数",
    "评价_总被引次数",
    "评价_顶级头衔标识",
    "评价_i10指数",
    "评价_TotalScore",
    "成功访问主页",
    "主页访问方式",
    "主页访问失败原因",
    "专家姓名验证",
    "姓名验证依据",
    "独立生存状态核验",
    "生存状态核验依据",
    "领域关联依据",
    "信息来源",
]

DELIVERY_COLUMNS = CORE_DELIVERY_COLUMNS + SUPPORTING_DELIVERY_COLUMNS

RANKING_COLUMNS = [
    "成功访问主页",
    "主页访问方式",
    "主页访问失败原因",
    "专家姓名验证",
    "姓名验证依据",
    "评价_H指数",
    "评价_总被引次数",
    "评价_顶级头衔标识",
    "评价_i10指数",
    "评价_TotalScore",
]

TOTAL_DATA_SOURCE_SUMMARY = "总数据来源：公开学术数据库、公开网页与专家主页、智能体结构化整理。"
TOTAL_DATA_SOURCE_WITH_EXTRA_SUMMARY = "总数据来源：公开学术数据库、公开网页与专家主页、用户补充数据源、智能体结构化整理。"

INTERNAL_ENRICHMENT_COLUMNS = [
    "OpenAlex匹配姓名",
    "OpenAlex作者ID",
    "OpenAlex主题",
    "S2匹配姓名",
    "S2作者ID",
    "S2主页",
    "S2代表论文",
]

GENERATED_METRIC_COLUMNS = RANKING_COLUMNS + INTERNAL_ENRICHMENT_COLUMNS + ["i10指数", "总被引次数"]

WORKPLACE_COLUMN_INDEX = 5

GREATER_CHINA_AFFILIATION_MARKERS = [
    "(CN)", "(CHN)", "(HK)", "(HKG)", "(MO)", "(MAC)", "(TW)", "(TWN)",
    "China", "Chinese Academy", "Chinese University", "Tsinghua", "Peking University",
    "Beijing University", "Zhejiang University", "Fudan University", "Shanghai Jiao Tong",
    "University of Science and Technology of China", "Harbin Institute of Technology",
    "Huazhong University", "Nanjing University", "Sun Yat-sen University",
    "Xi'an Jiaotong", "Tongji University", "Wuhan University", "Hong Kong",
    "Macau", "Macao", "Taiwan",
    "中国", "中华", "中科院", "中国科学院", "清华", "北京大学", "北大", "浙江大学",
    "复旦", "上海交通", "上海交大", "哈尔滨工业", "哈工大", "华中科技", "南京大学",
    "中山大学", "西安交通", "西安交大", "同济", "武汉大学", "香港", "澳门", "台湾",
]

# 仅用于剔除“明显跨领域错配”，不用于判断专家排名或替代领域证据核验。
# 强冲突词按组组织；目标领域没有任何正向证据、且命中任一强冲突组时才剔除。
DOMAIN_RELEVANCE_PROFILES = {
    "农业": {
        "aliases": ["农业", "农学", "农业科学", "agriculture", "agricultural", "agronomy"],
        "positive": [
            "农业", "农学", "作物", "育种", "种质", "农艺", "植物科学", "植物遗传",
            "园艺", "土壤", "农机", "农业工程", "农业信息", "精准农业", "智慧农业",
            "畜牧", "养殖", "动物科学", "兽医", "水产", "林业", "粮食", "种子",
            "crop", "agriculture", "agricultural", "agronomy", "plant breeding",
            "plant science", "horticulture", "soil", "livestock", "animal science",
            "veterinary", "aquaculture", "forestry", "farming", "seed science",
        ],
        "conflicts": [
            ["精神病", "精神医学", "精神卫生", "精神分裂", "psychiatry", "psychiatric", "schizophrenia", "psychosis"],
            ["量子场论", "粒子物理", "高能物理", "quantum field theory", "particle physics", "high energy physics"],
            ["天体物理", "宇宙学", "黑洞", "astrophysics", "cosmology", "black hole"],
            ["心脏外科", "神经外科", "整形外科", "cardiac surgery", "neurosurgery", "plastic surgery"],
        ],
    },
    "计算机": {
        "aliases": ["计算机", "信息技术", "人工智能", "computer science", "artificial intelligence"],
        "positive": [
            "计算机", "人工智能", "机器学习", "深度学习", "软件", "算法", "数据库",
            "网络安全", "计算机视觉", "自然语言处理", "信息检索", "机器人",
            "computer", "artificial intelligence", "machine learning", "deep learning",
            "software", "algorithm", "database", "cybersecurity", "computer vision",
            "natural language processing", "information retrieval", "robotics",
        ],
        "conflicts": [
            ["精神病", "精神医学", "精神分裂", "psychiatry", "schizophrenia"],
            ["作物育种", "种质资源", "畜牧育种", "crop breeding", "germplasm", "livestock breeding"],
            ["心脏外科", "神经外科", "cardiac surgery", "neurosurgery"],
        ],
    },
    "物理": {
        "aliases": ["物理", "physics"],
        "positive": [
            "物理", "量子", "凝聚态", "粒子", "光学", "光子", "等离子体", "宇宙学",
            "引力", "材料物理", "physics", "quantum", "condensed matter", "particle",
            "optics", "photonics", "plasma", "cosmology", "gravitation",
        ],
        "conflicts": [
            ["精神病", "精神医学", "精神分裂", "psychiatry", "schizophrenia"],
            ["作物育种", "种质资源", "畜牧育种", "crop breeding", "germplasm", "livestock breeding"],
            ["临床肿瘤", "心脏外科", "神经外科", "clinical oncology", "cardiac surgery", "neurosurgery"],
        ],
    },
    "医学": {
        "aliases": ["医学", "临床", "医疗", "medicine", "medical", "clinical"],
        "positive": [
            "医学", "临床", "疾病", "诊断", "治疗", "药物", "患者", "医院", "肿瘤",
            "心血管", "神经", "公共卫生", "medicine", "medical", "clinical", "disease",
            "diagnosis", "therapy", "drug", "patient", "hospital", "oncology",
        ],
        "conflicts": [
            ["量子场论", "粒子物理", "高能物理", "quantum field theory", "particle physics", "high energy physics"],
            ["天体物理", "宇宙学", "黑洞", "astrophysics", "cosmology", "black hole"],
            ["作物育种", "种质资源", "农机", "crop breeding", "germplasm", "agricultural machinery"],
        ],
    },
    "生物": {
        "aliases": ["生物", "生命科学", "biology", "biological science", "life science"],
        "positive": [
            "生物", "生命科学", "基因", "蛋白质", "细胞", "微生物", "生态", "免疫",
            "神经生物", "生物信息", "组学", "biology", "biological", "gene", "genome",
            "protein", "cell", "microbiology", "ecology", "immunology", "bioinformatics",
        ],
        "conflicts": [
            ["量子场论", "粒子物理", "高能物理", "quantum field theory", "particle physics", "high energy physics"],
            ["天体物理", "宇宙学", "黑洞", "astrophysics", "cosmology", "black hole"],
            ["软件工程", "数据库系统", "计算机网络", "software engineering", "database systems", "computer networks"],
        ],
    },
}

TITLE_BONUS_RULES = [
    ("院士", 30),
    ("Academician", 30),
    ("National Academy", 30),
    ("Royal Society", 30),
    ("IEEE Fellow", 30),
    ("ACM Fellow", 30),
    ("AAAI Fellow", 30),
    ("Fellow", 20),
    ("图灵奖", 30),
    ("Turing Award", 30),
    ("诺贝尔", 30),
    ("Nobel", 30),
    ("菲尔兹", 30),
    ("Fields Medal", 30),
]

def write_empty_expert_table(path: str) -> None:
    """
    写出空专家表，确保节点即使遇到解析失败也能产出可下载的 Excel 文件。
    """
    write_expert_excel(pd.DataFrame(columns=EXPERT_TABLE_HEADERS + RANKING_COLUMNS), path)


def _clean_text(value: object, fallback: str = "暂无公开信息") -> str:
    try:
        if pd.isna(value):
            return fallback
    except (TypeError, ValueError):
        pass
    text = str(value or "").strip()
    if not text or text.lower() in {"nan", "none"}:
        return fallback
    return text


def _append_information_source(row_data: dict, url: object) -> None:
    clean_url = str(url or "").strip()
    if not clean_url.startswith(("http://", "https://")):
        return

    current = str(row_data.get("信息来源", "") or "")
    urls = re.findall(r"https?://[^\s；;,，)）]+", current)
    if clean_url not in urls:
        urls.append(clean_url)
    row_data["信息来源"] = "；".join(dict.fromkeys(urls))


def remove_non_expert_rows(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty or "专家姓名" not in df.columns:
        return df

    names = df["专家姓名"].fillna("").astype(str).str.strip()
    valid = (
        names.ne("")
        & ~names.str.startswith("总数据来源：")
        & ~names.str.lower().isin({"nan", "none", "专家姓名", "姓名"})
    )
    return df.loc[valid].copy()


def query_domain_parts(query: str) -> tuple[str, str]:
    text = str(query or "").strip()
    match = re.search(r"(.+?)领域下的(.+?)方向顶级专家信息", text)
    if match:
        return match.group(1).strip(), match.group(2).strip()
    return text or "暂无公开信息", "暂无公开信息"


def normalize_subdomain_membership(value: object) -> str:
    """规范化专家所属细分领域，保留一位专家的多个真实归属。"""
    text = str(value or "").strip()
    if not text or _is_missing_or_pending(text):
        return "暂无公开信息"

    domains = []
    for item in re.split(r"[；;、|/\n]+", text):
        clean = re.sub(r"^(?:细分领域|所属细分领域)\s*[:：]\s*", "", item.strip())
        clean = re.sub(r"\s*(?:方向|领域)$", "", clean).strip()
        if clean and not _is_missing_or_pending(clean) and clean not in domains:
            domains.append(clean)
    return "；".join(domains) if domains else "暂无公开信息"


def _split_collaboration(value: object) -> tuple[str, str]:
    text = _clean_text(value)
    if text == "暂无公开信息":
        return text, text

    scholars = []
    units = []
    for segment in re.split(r"[；;]\s*", text):
        segment = segment.strip()
        if not segment:
            continue
        segment = re.sub(r"^(?:强合作|中合作|弱合作|合作线索|合作学者)\s*[:：]\s*", "", segment)
        evidence_free = re.split(r"[，,]\s*(?:合作依据|来源|证据)\s*[:：]", segment)[0].strip()
        parse_text = re.sub(r"\((?:CN|CHN|HK|HKG|MO|MAC|TW|TWN)\)", "", evidence_free, flags=re.IGNORECASE)
        pairs = re.findall(r"([^，,、；;（）()：:]{2,80})[（(]([^）)]+)[）)]", parse_text)
        if pairs:
            for scholar, inside in pairs:
                scholar = scholar.strip(" ，,、")
                if scholar and not re.search(
                    r"大学|学院|研究所|科学院|实验室|中心|公司|University|Institute|Academy",
                    scholar,
                    flags=re.IGNORECASE,
                ):
                    scholars.append(scholar)
                for unit in re.split(r"[，,、]\s*", inside):
                    unit = unit.strip()
                    if re.search(
                        r"大学|学院|研究所|科学院|实验室|中心|公司|University|Institute|Academy|\(CN\)",
                        unit,
                        flags=re.IGNORECASE,
                    ):
                        units.append(unit)
        elif re.search(r"大学|学院|研究所|科学院|实验室|中心|公司|University|Institute|Academy", evidence_free, flags=re.IGNORECASE):
            units.append(evidence_free)

    return (
        "；".join(dict.fromkeys(scholars)) or "暂无公开信息",
        "；".join(dict.fromkeys(units)) or "暂无公开信息",
    )


def _row_source_urls(row: pd.Series) -> str:
    urls = []
    for column in ["信息来源", "个人主页"]:
        for url in re.findall(r"https?://[^\s；;,，)）]+", str(row.get(column, ""))):
            clean_url = url.rstrip("。.")
            if clean_url not in urls:
                urls.append(clean_url)
    return "；".join(urls) if urls else "暂无可追溯网页来源"


def selection_basis(row: pd.Series) -> str:
    """基于已有证据生成保守的入选理由，不推断未经证实的头衔。"""
    explicit = _clean_text(row.get("入选依据"))
    if explicit != "暂无公开信息":
        return explicit

    title_evidence = _clean_text(row.get("评价_顶级头衔标识"))
    if title_evidence not in {"暂无公开信息", "无"}:
        return f"因具有可验证的权威头衔、奖项或行业组织身份入选：{title_evidence}"

    evidence_text = "；".join(
        str(row.get(column, "") or "")
        for column in ["职位", "主要成果", "工作经历"]
    )
    high_value_patterns = [
        r"[^；。]{0,80}(?:美国国家科学院|National Academy|院士|Academician)[^；。]{0,80}",
        r"[^；。]{0,80}(?:CIGR|IEEE Fellow|ACM Fellow|AAAI Fellow|Royal Society)[^；。]{0,80}",
        r"[^；。]{0,80}(?:世界粮食奖|World Food Prize|诺贝尔|Nobel|图灵奖|Turing Award|菲尔兹|Fields Medal)[^；。]{0,80}",
        r"[^；。]{0,80}(?:全球前2%|斯坦福2%|高被引学者|Highly Cited Researcher)[^；。]{0,80}",
    ]
    evidence_hits = []
    for pattern in high_value_patterns:
        evidence_hits.extend(re.findall(pattern, evidence_text, flags=re.IGNORECASE))
    evidence_hits = [hit.strip(" ；。") for hit in evidence_hits if hit.strip(" ；。")]
    if evidence_hits:
        return "因公开证据显示其具有重要头衔、奖项、榜单或行业组织身份入选：" + "；".join(
            dict.fromkeys(evidence_hits[:3])
        )

    h_index = parse_metric_number(row.get("H指数", 0))
    citations = estimate_citations(row)
    metrics = []
    if h_index >= 30:
        metrics.append(f"H指数约{int(h_index)}")
    if citations >= 10000:
        metrics.append(f"总被引次数约{int(citations)}")
    if metrics:
        return "因学术影响力指标突出入选：" + "；".join(metrics)

    relevance = _clean_text(row.get("领域关联依据"), _clean_text(row.get("主要成果")))
    if relevance != "暂无公开信息":
        return f"因其代表成果与目标细分领域直接相关而入选：{relevance}"
    return "暂无公开信息"


def build_delivery_dataframe(df: pd.DataFrame, query: str = "") -> pd.DataFrame:
    main_domain, sub_domain = query_domain_parts(query)
    rows = []
    for _, row in remove_non_expert_rows(df).iterrows():
        scholars, units = _split_collaboration(row.get("国内合作学者与单位", ""))
        rows.append(
            {
                "姓名": _clean_text(row.get("专家姓名")),
                "国籍": _clean_text(row.get("国籍")),
                "所在工作机构": _clean_text(row.get("工作单位")),
                "职位": _clean_text(row.get("职位")),
                "学科领域": _clean_text(row.get("学科领域"), main_domain),
                "细分领域": normalize_subdomain_membership(
                    _clean_text(row.get("细分领域"), sub_domain)
                ),
                "领域关联依据": _clean_text(
                    row.get("领域关联依据"),
                    _clean_text(row.get("研究兴趣")),
                ),
                "工作经历": _clean_text(row.get("工作经历")),
                "主要成果": _clean_text(row.get("主要成果")),
                "联系方式": _clean_text(row.get("邮箱/电话")),
                "国内合作学者": _clean_text(row.get("国内合作学者"), scholars),
                "国内合作单位": _clean_text(row.get("国内合作单位"), units),
                "入选依据": selection_basis(row),
                "个人主页": _clean_text(row.get("个人主页")),
                "研究兴趣": _clean_text(row.get("研究兴趣")),
                "教育背景": _clean_text(row.get("教育背景")),
                "性别": _clean_text(row.get("性别")),
                "语言": _clean_text(row.get("语言")),
                "H指数": _clean_text(row.get("H指数")),
                "G指数": _clean_text(row.get("G指数")),
                "i10指数": _clean_text(row.get("i10指数")),
                "总被引次数": _clean_text(row.get("总被引次数")),
                "评价_H指数": _clean_text(row.get("评价_H指数")),
                "评价_总被引次数": _clean_text(row.get("评价_总被引次数")),
                "评价_顶级头衔标识": _clean_text(row.get("评价_顶级头衔标识")),
                "评价_i10指数": _clean_text(row.get("评价_i10指数")),
                "评价_TotalScore": _clean_text(row.get("评价_TotalScore")),
                "成功访问主页": _clean_text(row.get("成功访问主页"), "未尝试"),
                "主页访问方式": _clean_text(row.get("主页访问方式"), "未尝试"),
                "主页访问失败原因": _clean_text(row.get("主页访问失败原因")),
                "专家姓名验证": _clean_text(row.get("专家姓名验证"), "待复核"),
                "姓名验证依据": _clean_text(row.get("姓名验证依据")),
                "独立生存状态核验": _clean_text(row.get("独立生存状态核验"), "未执行"),
                "生存状态核验依据": _clean_text(row.get("生存状态核验依据")),
                "信息来源": _row_source_urls(row),
            }
        )
    return pd.DataFrame(rows, columns=DELIVERY_COLUMNS)


def build_explanation_dataframe(
    query: str = "",
    extra_data_source_urls: list[str] | None = None,
    supplemental_document_names: list[str] | None = None,
    data_source_summary_text: str | None = None,
) -> pd.DataFrame:
    rows = [
        ("运行信息", "检索任务", query or "暂无公开信息"),
        ("抓取逻辑", "候选发现", "从公开学术数据库、官方人才名单、奖项官网、协会名录和网页检索结果发现候选专家。"),
        ("抓取逻辑", "领域关联核验", "依据研究主题、代表成果、项目、专利或奖项，判断专家与目标细分领域的直接关联。"),
        ("抓取逻辑", "身份与状态核验", "通过 OpenAlex、Semantic Scholar 和官方个人主页交叉验证身份；生存状态由独立网页核验层查找讣告、逝世公告等证据，只有确认已故时才剔除。"),
        ("抓取逻辑", "信息补全", "优先访问个人主页补充机构、职位、工作经历、成果、联系方式与国内合作证据。"),
        ("抓取逻辑", "中文交付处理", "最终甲方主表先执行常用术语规范化，再对英文叙述进行批量中文翻译，并生成“中文输出检查”工作表记录残留内容。"),
        ("筛选标准", "入选要求", "目标细分领域直接相关；身份可交叉验证；具有高影响力指标、重要成果、奖项、协会身份或行业贡献。"),
        ("筛选标准", "排除规则", "排除明确已故、明显同名错配、领域关联不足、缺乏可追溯来源或信息明显不可靠的候选人。"),
        ("评分逻辑", "综合评价得分", "TotalScore = H指数 + log10(总被引次数) × 15 + i10指数 × 0.2 + 顶级头衔加分。"),
        ("评分逻辑", "使用说明", "综合评价得分用于排序，不替代身份核验、领域关联判断与人工复核。"),
        ("数据来源", "总览", data_source_summary_text or TOTAL_DATA_SOURCE_SUMMARY),
    ]
    for source in source_records_for_query(query, extra_data_source_urls):
        rows.append(
            (
                "数据来源",
                str(source.get("name", "公开数据源")),
                f"{source.get('url', '')} | 用途：{source.get('purpose', '')}",
            )
        )
    for name in supplemental_document_names or []:
        rows.append(
            (
                "数据来源",
                "用户上传补充文件",
                f"{name} | 用途：候选专家发现与公开信息补充核验",
            )
        )
    return pd.DataFrame(rows, columns=["类别", "项目", "内容/网址"])


def build_run_summary_dataframe(df: pd.DataFrame, query: str = "") -> pd.DataFrame:
    def count_value(column: str, value: str) -> int:
        if column not in df.columns:
            return 0
        return int(df[column].fillna("").astype(str).str.strip().eq(value).sum())

    rows = [
        ("生成时间", time.strftime("%Y-%m-%d %H:%M:%S"), "本地系统时间"),
        ("检索任务", query or "暂无公开信息", "用户输入的大领域、细分领域和梯队条件"),
        ("最终输出专家数", len(df), "已完成清洗、去重和明确已故专家过滤"),
        ("姓名已验证", count_value("专家姓名验证", "已验证"), "通过多个公开来源交叉验证"),
        ("姓名待复核", count_value("专家姓名验证", "待复核"), "建议在正式使用前人工复核"),
        ("成功访问主页", count_value("成功访问主页", "是"), "主页正文可访问且包含专家姓名"),
        ("OpenCLI成功访问主页", count_value("主页访问方式", "OpenCLI"), "普通 HTTP 失败后由真实浏览器渲染读取成功"),
        ("主页访问失败", count_value("成功访问主页", "否"), "包括 403、404、超时、SSL 或身份不匹配"),
        ("主页未尝试", count_value("成功访问主页", "未尝试"), "受补全上限或缺少有效主页影响"),
        ("疑似已故待复核", count_value("独立生存状态核验", "疑似已故待复核"), "未自动剔除，建议人工复核"),
        ("未发现死亡证据", count_value("独立生存状态核验", "未发现死亡证据"), "独立网页检索未发现明确死亡证据"),
        ("生存状态核验未执行", count_value("独立生存状态核验", "未执行"), "受配置、数量上限或缺少姓名影响"),
    ]
    return pd.DataFrame(rows, columns=["项目", "结果", "说明"])


def _style_excel_sheet(worksheet, header_row: int = 1, freeze_cell: str = "A2") -> None:
    header_fill = PatternFill("solid", fgColor="2F6B4F")
    header_font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
    body_font = Font(name="Microsoft YaHei", size=10)
    thin = Side(style="thin", color="D9E2DD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    worksheet.freeze_panes = freeze_cell
    worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(worksheet.max_column)}{worksheet.max_row}"
    for cell in worksheet[header_row]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    for row in worksheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = body_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = border
            if isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                cell.hyperlink = cell.value.split(" | ", 1)[0]
                cell.style = "Hyperlink"

    for index in range(1, worksheet.max_column + 1):
        letter = get_column_letter(index)
        max_length = max(
            [len(str(worksheet.cell(row=row, column=index).value or "")) for row in range(1, min(worksheet.max_row, 30) + 1)]
            or [8]
        )
        worksheet.column_dimensions[letter].width = min(max(max_length * 1.15, 12), 42)


def write_expert_excel(
    df: pd.DataFrame,
    path: str,
    data_source_summary: str | None = None,
    query: str = "",
    extra_data_source_urls: list[str] | None = None,
    supplemental_document_names: list[str] | None = None,
    translate_delivery: bool = False,
    final_delivery_only: bool = False,
) -> None:
    """
    写出专家 Excel。

    中间批次默认保留评价验证等内部工作表；最终总名单使用参考文件的
    单工作表平铺样式，只输出甲方交付字段。
    """
    clean_df = remove_non_expert_rows(df)
    delivery_df = build_delivery_dataframe(clean_df, query)
    delivery_df, chinese_audit_df = translate_delivery_dataframe(
        delivery_df,
        enable_llm=translate_delivery,
    )
    explanation_df = build_explanation_dataframe(
        query,
        extra_data_source_urls,
        supplemental_document_names,
        data_source_summary,
    )
    run_summary_df = build_run_summary_dataframe(clean_df, query)
    residual_count = int(chinese_audit_df["检查结果"].eq("存在英文残留").sum())
    run_summary_df.loc[len(run_summary_df)] = [
        "中文输出检查残留项",
        residual_count,
        "详见“中文输出检查”工作表；正式交付前建议复核专名与未翻译叙述。",
    ]
    main_domain, _ = query_domain_parts(query)

    if final_delivery_only:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            delivery_df.to_excel(writer, index=False, sheet_name="Sheet1")
            worksheet = writer.book["Sheet1"]
            _style_excel_sheet(worksheet, header_row=1, freeze_cell="A2")
            summary_row = len(delivery_df) + 3
            worksheet.merge_cells(
                start_row=summary_row,
                start_column=1,
                end_row=summary_row,
                end_column=len(DELIVERY_COLUMNS),
            )
            summary_cell = worksheet.cell(
                row=summary_row,
                column=1,
                value=data_source_summary or TOTAL_DATA_SOURCE_SUMMARY,
            )
            summary_cell.alignment = Alignment(wrap_text=True, vertical="top")
            summary_cell.font = Font(name="Microsoft YaHei", size=10)
            worksheet.row_dimensions[summary_row].height = 180
        return

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        delivery_df.to_excel(writer, index=False, sheet_name="专家人才信息", startrow=1)
        clean_df.to_excel(writer, index=False, sheet_name="评价与验证")
        explanation_df.to_excel(writer, index=False, sheet_name="检索与评分说明")
        run_summary_df.to_excel(writer, index=False, sheet_name="运行摘要")
        chinese_audit_df.to_excel(writer, index=False, sheet_name="中文输出检查")

        main_sheet = writer.book["专家人才信息"]
        main_sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(DELIVERY_COLUMNS))
        title_cell = main_sheet.cell(row=1, column=1, value=f"{main_domain}关键人才信息")
        title_cell.font = Font(name="Microsoft YaHei", size=16, bold=True, color="1F2937")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        main_sheet.row_dimensions[1].height = 28
        _style_excel_sheet(main_sheet, header_row=2, freeze_cell="A3")
        _style_excel_sheet(writer.book["评价与验证"], header_row=1, freeze_cell="A2")
        _style_excel_sheet(writer.book["检索与评分说明"], header_row=1, freeze_cell="A2")
        _style_excel_sheet(writer.book["运行摘要"], header_row=1, freeze_cell="A2")
        _style_excel_sheet(writer.book["中文输出检查"], header_row=1, freeze_cell="A2")


def data_source_summary(
    extra_data_source_urls: list[str] | None = None,
    query: str = "",
    supplemental_document_names: list[str] | None = None,
) -> str:
    """
    生成可审计的具体数据来源说明。

    “本轮系统来源”只列出代码实际调用或读取的网页；Google Scholar 与
    Scopus 当前未被系统直接批量抓取，因此单独标注为人工复核渠道。
    """
    source_lines = []
    for record in source_records_for_query(query, extra_data_source_urls):
        name = str(record.get("name", "")).strip()
        url = str(record.get("url", "")).strip()
        purpose = str(record.get("purpose", "")).strip()
        if name and url:
            source_lines.append(f"- {name}：{url}（{purpose}）")

    review_lines = [
        "- Google Scholar：https://scholar.google.com/（建议人工复核论文、引用和学者身份；系统当前不直接批量抓取）",
        "- Scopus：https://www.scopus.com/（建议人工复核作者指标与机构信息；系统当前不直接批量抓取）",
    ]
    document_lines = [
        f"- 用户上传补充文件：{name}（用于候选发现与信息补充；关键事实仍需公开来源交叉验证）"
        for name in supplemental_document_names or []
    ]
    return (
        "具体数据来源说明：\n"
        "一、本轮系统实际使用或读取的来源：\n"
        + "\n".join(source_lines)
        + ("\n二、用户上传补充文件：\n" + "\n".join(document_lines) if document_lines else "")
        + (("\n三、建议人工复核渠道：\n") if document_lines else "\n二、建议人工复核渠道：\n")
        + "\n".join(review_lines)
    )

def parse_metric_number(value) -> float:
    """
    从 H/i10/引用数字段中提取数值，兼容“约85”“50,000”“暂无公开信息”等文本。
    """
    if pd.isna(value):
        return 0.0

    text = str(value).replace(",", "")
    match = re.search(r"\d+(?:\.\d+)?", text)
    if not match:
        return 0.0

    return float(match.group(0))

def estimate_citations(row: pd.Series) -> float:
    """
    优先读取明确总引用列；缺失时从主要成果文本中提取“引用/被引”附近的数值。
    """
    for column in ["总被引次数", "总引用数", "Total Citations", "Citations", "引用次数", "被引次数"]:
        if column in row.index:
            value = parse_metric_number(row.get(column))
            if value > 0:
                return value

    evidence_text = " ".join(
        str(row.get(column, ""))
        for column in ["主要成果", "研究兴趣", "职位", "教育背景"]
        if column in row.index
    )
    number_pattern = r"(\d+(?:,\d{3})*(?:\.\d+)?)"
    unit_pattern = r"(万|千|[kKmM]|million|thousand)?"
    number_before_matches = re.findall(
        rf"{number_pattern}\s*{unit_pattern}\s*(?:次)?\s*(?:引用|被引|citations?|cited)",
        evidence_text,
        flags=re.IGNORECASE,
    )
    number_after_matches = re.findall(
        rf"(?:引用|被引|被引用|citations?|cited)\s*"
        rf"(?:约|超|超过|逾|over|more\s+than)?\s*"
        rf"{number_pattern}\s*{unit_pattern}\s*(?:次)?",
        evidence_text,
        flags=re.IGNORECASE,
    )

    def scaled_value(match: tuple[str, str]) -> float:
        value = parse_metric_number(match[0])
        unit = str(match[1] or "").lower()
        if unit in {"万"}:
            return value * 10000
        if unit in {"千", "k", "thousand"}:
            return value * 1000
        if unit in {"m", "million"}:
            return value * 1000000
        return value

    citation_values = [
        scaled_value(match)
        for match in number_before_matches + number_after_matches
    ]
    return max(citation_values) if citation_values else 0.0

def get_i10_metric(row: pd.Series) -> float:
    """
    优先使用 i10 指数；如果没有该列，临时使用现有 G 指数列作为产出广度代理。
    """
    for column in ["i10指数", "i10-index", "I10指数", "I10"]:
        if column in row.index:
            value = parse_metric_number(row.get(column))
            if value > 0:
                return value

    return parse_metric_number(row.get("G指数", 0))

def title_bonus(row: pd.Series) -> tuple[int, str]:
    evidence_text = " ".join(
        str(row.get(column, ""))
        for column in ["职位", "主要成果", "教育背景", "工作单位", "入选依据"]
        if column in row.index
    )

    matched = []
    best_bonus = 0
    for keyword, bonus in TITLE_BONUS_RULES:
        if keyword.lower() in evidence_text.lower():
            matched.append(keyword)
            best_bonus = max(best_bonus, bonus)

    if any(keyword.endswith("Fellow") and keyword != "Fellow" for keyword in matched):
        matched = [keyword for keyword in matched if keyword != "Fellow"]

    return best_bonus, "、".join(dict.fromkeys(matched)) if matched else "无"


def normalize_expert_name_for_dedupe(name: object) -> str:
    """
    近似姓名去重：忽略 Sir/Prof 等头衔和单字母中间名。
    只有在剩余有效词至少两个时才启用，避免只按姓氏误合并。
    """
    text = str(name).lower()
    text = re.sub(r"\b(sir|prof|professor|dr|phd|md)\b", " ", text)
    tokens = re.findall(r"[a-zA-Z\u4e00-\u9fff]+", text)
    meaningful_tokens = [token for token in tokens if len(token) > 1]
    if len(meaningful_tokens) >= 2:
        return " ".join(meaningful_tokens)
    return " ".join(tokens)


def _row_completeness(row: pd.Series) -> int:
    score = 0
    for column in ["个人主页", "邮箱/电话", "教育背景", "国内合作学者与单位", "OpenAlex作者ID", "S2作者ID"]:
        value = str(row.get(column, ""))
        if value and not re.search(r"暂无|未知|公开信息|nan|None", value, flags=re.IGNORECASE):
            score += 1
    return score


def _is_missing_or_pending(value: object) -> bool:
    text = str(value or "").strip()
    return not text or bool(
        re.search(
            r"暂无|未知|公开信息|待补充|待核验|未核实|候选未核实|nan|None",
            text,
            flags=re.IGNORECASE,
        )
    )


def _merge_text_values(left: object, right: object, *, combine: bool = False) -> str:
    left_text = str(left or "").strip()
    right_text = str(right or "").strip()
    if _is_missing_or_pending(left_text) and not _is_missing_or_pending(right_text):
        return right_text
    if _is_missing_or_pending(right_text):
        return left_text
    if not left_text:
        return right_text
    if not right_text or left_text == right_text:
        return left_text
    if combine:
        parts = [
            part.strip()
            for value in [left_text, right_text]
            for part in re.split(r"[；;]\s*", value)
            if part.strip() and not _is_missing_or_pending(part)
        ]
        return "；".join(dict.fromkeys(parts))
    return right_text if len(right_text) > len(left_text) else left_text


def _merge_duplicate_group(group: pd.DataFrame) -> pd.Series:
    base = group.iloc[0].copy()
    combine_columns = {
        "主要成果",
        "国内合作学者与单位",
        "入选依据",
        "领域关联依据",
        "信息来源",
        "细分领域",
        "姓名验证依据",
        "生存状态核验依据",
    }
    numeric_max_columns = {
        "H指数",
        "i10指数",
        "总被引次数",
        "评价_H指数",
        "评价_总被引次数",
        "评价_i10指数",
        "评价_TotalScore",
    }
    for _, candidate in group.iloc[1:].iterrows():
        for column in group.columns:
            if column.startswith("_"):
                continue
            if column in numeric_max_columns:
                left_number = parse_metric_number(base.get(column, 0))
                right_number = parse_metric_number(candidate.get(column, 0))
                if right_number > left_number:
                    base[column] = candidate.get(column)
                continue
            base[column] = _merge_text_values(
                base.get(column, ""),
                candidate.get(column, ""),
                combine=column in combine_columns,
            )
    if "细分领域" in base.index:
        base["细分领域"] = normalize_subdomain_membership(base.get("细分领域", ""))
    return base


def drop_near_duplicate_experts(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty or "专家姓名" not in df.columns:
        return df

    work_df = df.copy()
    work_df["_dedupe_name_key"] = work_df["专家姓名"].apply(normalize_expert_name_for_dedupe)
    work_df["_dedupe_completeness"] = work_df.apply(_row_completeness, axis=1)
    sort_columns = [
        column for column in ["评价_TotalScore", "评价_总被引次数", "_dedupe_completeness"]
        if column in work_df.columns
    ]
    if sort_columns:
        work_df = work_df.sort_values(sort_columns, ascending=False, na_position="last")

    merged_rows = [
        _merge_duplicate_group(group)
        for _, group in work_df.groupby("_dedupe_name_key", sort=False, dropna=False)
    ]
    merged_df = pd.DataFrame(merged_rows)
    return merged_df.drop(columns=["_dedupe_name_key", "_dedupe_completeness"], errors="ignore")


def _workplace_column(df: pd.DataFrame) -> str | None:
    for column in ["工作单位", "宸ヤ綔鍗曚綅", "Institution", "Affiliation"]:
        if column in df.columns:
            return column
    if len(df.columns) > WORKPLACE_COLUMN_INDEX:
        return df.columns[WORKPLACE_COLUMN_INDEX]
    return None


def filter_foreign_experts(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df

    workplace_column = _workplace_column(df)
    if not workplace_column:
        return df

    def is_foreign_affiliation(value: object) -> bool:
        text = str(value or "")
        if not text.strip():
            return True
        return not any(marker.lower() in text.lower() for marker in GREATER_CHINA_AFFILIATION_MARKERS)

    filtered_df = df[df[workplace_column].apply(is_foreign_affiliation)].copy()
    removed_count = len(df) - len(filtered_df)
    if removed_count:
        safe_print(f"[Foreign expert filter] Removed {removed_count} China/Greater China affiliated rows.")
    return filtered_df


def is_metric_match_trusted(row_data: dict, metrics: dict[str, object], source: str = "") -> bool:
    """
    用通用可信度规则过滤数据库错配，不再依赖手写领域词典。
    """
    current_h = parse_metric_number(row_data.get("H指数", 0))
    current_citations = estimate_citations(pd.Series(row_data))
    metric_h = parse_metric_number(metrics.get("h_index"))
    metric_citations = parse_metric_number(
        metrics.get("cited_by_count")
        if "cited_by_count" in metrics
        else metrics.get("citation_count")
    )

    if current_h >= 20 and metric_h > 0 and metric_h < current_h * 0.65:
        return False
    if current_h >= 20 and metric_citations > 0 and metric_citations < max(100, current_h * 30):
        return False
    if current_h >= 10 and 0 < metric_h <= 5 and metric_citations < 500:
        return False
    if current_citations >= 1000 and 0 < metric_citations < current_citations * 0.2:
        return False

    # 如果数据库只返回极低指标，而模型原始表中没有可靠指标，也不要把它当作强证据写入。
    if current_h == 0 and current_citations == 0 and 0 < metric_h <= 2 and metric_citations < 100:
        return False

    return True

def clean_contact_fields(df: pd.DataFrame) -> pd.DataFrame:
    """
    清洗模型生成的伪具体主页/邮箱，如“某大学官方主页”“某大学官方邮箱”。
    """
    if "个人主页" in df.columns:
        df["个人主页"] = df["个人主页"].apply(
            lambda value: value
            if "http://" in str(value) or "https://" in str(value) or "暂无公开信息" in str(value)
            else "暂无公开信息"
        )

    if "邮箱/电话" in df.columns:
        df["邮箱/电话"] = df["邮箱/电话"].apply(
            lambda value: re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", str(value)).group(0)
            if re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", str(value))
            else "暂无公开信息"
        )

    return df


def _public_value(value: object) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    if re.search(r"暂无|未知|公开信息|nan|None|未公开|无公开", text, flags=re.IGNORECASE):
        return ""
    return text


def _same_verified_name(left: object, right: object) -> bool:
    left_key = normalize_expert_name_for_dedupe(left)
    right_key = normalize_expert_name_for_dedupe(right)
    if not left_key or not right_key:
        return False

    if left_key == right_key:
        return True

    left_compact = left_key.replace(" ", "")
    right_compact = right_key.replace(" ", "")
    if re.search(r"[\u4e00-\u9fff]", left_compact + right_compact):
        return left_compact == right_compact

    left_tokens = set(left_key.split())
    right_tokens = set(right_key.split())
    return len(left_tokens & right_tokens) >= 2 and (
        left_tokens.issubset(right_tokens) or right_tokens.issubset(left_tokens)
    )


def verify_expert_name(row: pd.Series) -> tuple[str, str]:
    name = row.get("专家姓名", "")
    evidence = []
    suspicious = []

    openalex_name = _public_value(row.get("OpenAlex匹配姓名", ""))
    openalex_id = _public_value(row.get("OpenAlex作者ID", ""))
    if openalex_name or openalex_id:
        if not openalex_name or _same_verified_name(name, openalex_name):
            evidence.append(f"OpenAlex作者记录{f'({openalex_name})' if openalex_name else ''}")
        else:
            suspicious.append(f"OpenAlex疑似错配({openalex_name})")

    s2_name = _public_value(row.get("S2匹配姓名", ""))
    s2_id = _public_value(row.get("S2作者ID", ""))
    if s2_name or s2_id:
        if not s2_name or _same_verified_name(name, s2_name):
            evidence.append(f"Semantic Scholar作者记录{f'({s2_name})' if s2_name else ''}")
        else:
            suspicious.append(f"Semantic Scholar疑似错配({s2_name})")

    if str(row.get("成功访问主页", "")).strip() == "是":
        evidence.append("专家个人主页访问成功")

    if suspicious:
        return "待复核", "；".join(dict.fromkeys(suspicious + evidence))
    if len(evidence) >= 2:
        return "已验证", "；".join(dict.fromkeys(evidence))
    if evidence:
        return "待复核", f"当前仅有单一验证来源：{'；'.join(dict.fromkeys(evidence))}"
    return "待复核", "未在 OpenAlex、Semantic Scholar 或可访问个人主页中完成姓名验证"


def add_name_verification_columns(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        df["专家姓名验证"] = []
        df["姓名验证依据"] = []
        return df

    verification_rows = [
        {
            "专家姓名验证": status,
            "姓名验证依据": evidence,
        }
        for status, evidence in df.apply(verify_expert_name, axis=1)
    ]
    verification_df = pd.DataFrame(verification_rows)
    return pd.concat([df.reset_index(drop=True), verification_df], axis=1)


def enrich_openalex_metrics(df: pd.DataFrame, query: str = "") -> pd.DataFrame:
    """
    用 OpenAlex 按专家姓名补齐 H 指数、总被引次数、i10 指数等结构化指标。
    """
    if df.empty or "专家姓名" not in df.columns:
        return df

    max_row_errors = int(os.environ.get("OPENALEX_MAX_ROW_ERRORS", "8"))
    openalex_error_count = 0
    openalex_paused = False
    enriched_rows = []
    for _, row in df.iterrows():
        row_data = row.to_dict()
        name = str(row.get("专家姓名", "")).strip()
        metrics = {}
        if not openalex_paused and name and name not in {"暂无公开信息", "未知", "nan"}:
            try:
                metrics = get_openalex_author_metrics(name, query or "")
            except Exception as e:
                openalex_error_count += 1
                safe_print(f"[OpenAlex警告] {name} 指标回填失败: {e}")
                if openalex_error_count >= max_row_errors or "OpenAlex 暂停请求" in str(e):
                    openalex_paused = True
                    safe_print(
                        f"[OpenAlex限流保护] 本批指标回填已暂停，"
                        f"避免继续触发限流或连接错误。累计错误: {openalex_error_count}"
                    )

        if metrics:
            current_h = parse_metric_number(row_data.get("H指数", 0))
            current_i10 = get_i10_metric(pd.Series(row_data))
            current_citations = estimate_citations(pd.Series(row_data))
            openalex_h = parse_metric_number(metrics.get("h_index"))
            openalex_i10 = parse_metric_number(metrics.get("i10_index"))
            openalex_citations = parse_metric_number(metrics.get("cited_by_count"))
            metrics_trusted = is_metric_match_trusted(row_data, metrics, "openalex")

            if metrics_trusted and openalex_h and openalex_h >= current_h:
                row_data["H指数"] = metrics["h_index"]
            if metrics_trusted and openalex_i10 and openalex_i10 >= current_i10:
                row_data["i10指数"] = metrics["i10_index"]
            if metrics_trusted and openalex_citations and openalex_citations >= current_citations:
                row_data["总被引次数"] = metrics["cited_by_count"]

            if metrics_trusted:
                row_data["OpenAlex匹配姓名"] = metrics.get("openalex_name", "")
                row_data["OpenAlex作者ID"] = metrics.get("openalex_id", "")
                row_data["OpenAlex主题"] = metrics.get("topics", "")
                _append_information_source(row_data, metrics.get("openalex_id"))
            else:
                row_data["OpenAlex匹配姓名"] = ""
                row_data["OpenAlex作者ID"] = ""
                row_data["OpenAlex主题"] = ""
        else:
            row_data["OpenAlex匹配姓名"] = ""
            row_data["OpenAlex作者ID"] = ""
            row_data["OpenAlex主题"] = ""

        enriched_rows.append(row_data)

    return pd.DataFrame(enriched_rows)

def enrich_semantic_scholar_metrics(df: pd.DataFrame, query: str = "") -> pd.DataFrame:
    """
    按需用 Semantic Scholar 补充作者指标、主页和代表论文。
    默认只处理少量高价值行，避免 1 request/sec 限速拖慢整批任务。
    """
    if df.empty or "专家姓名" not in df.columns:
        return df

    max_rows = int(os.environ.get("SEMANTIC_SCHOLAR_MAX_ROWS", "2"))
    if max_rows <= 0:
        for column in ["S2匹配姓名", "S2作者ID", "S2主页", "S2代表论文"]:
            if column not in df.columns:
                df[column] = ""
        return df

    work_df = df.copy()
    for column in ["S2匹配姓名", "S2作者ID", "S2主页", "S2代表论文"]:
        if column not in work_df.columns:
            work_df[column] = ""

    priorities = []
    for index, row in work_df.iterrows():
        openalex_missing = not str(row.get("OpenAlex作者ID", "")).strip()
        citations = estimate_citations(row)
        h_index = parse_metric_number(row.get("H指数", 0))
        priority = 0
        if openalex_missing:
            priority += 1000
        if citations <= 0:
            priority += 800
        if h_index >= 50:
            priority += 400
        if not str(row.get("S2代表论文", "")).strip():
            priority += 100
        priority += min(h_index, 200)
        priorities.append((priority, index))

    selected_indices = [
        index for priority, index in sorted(priorities, reverse=True)
        if priority > 0
    ][:max_rows]
    selected_set = set(selected_indices)

    enriched_rows = []
    for index, row in work_df.iterrows():
        row_data = row.to_dict()
        if index not in selected_set:
            enriched_rows.append(row_data)
            continue

        name = str(row.get("专家姓名", "")).strip()
        metrics = {}
        if name and name not in {"暂无公开信息", "未知", "nan"}:
            try:
                metrics = get_semantic_scholar_author_metrics(name, query or "")
            except Exception as e:
                safe_print(f"[SemanticScholar警告] {name} 指标回填失败: {e}")

        if metrics:
            current_h = parse_metric_number(row_data.get("H指数", 0))
            current_citations = estimate_citations(pd.Series(row_data))
            s2_h = parse_metric_number(metrics.get("h_index"))
            s2_citations = parse_metric_number(metrics.get("citation_count"))

            metrics_trusted = is_metric_match_trusted(row_data, metrics, "semantic_scholar")

            if metrics_trusted and s2_h and s2_h >= current_h:
                row_data["H指数"] = metrics["h_index"]
            if metrics_trusted and s2_citations and s2_citations >= current_citations:
                row_data["总被引次数"] = metrics["citation_count"]

            homepage = metrics.get("homepage")
            if metrics_trusted and homepage and str(row_data.get("个人主页", "")) == "暂无公开信息":
                row_data["个人主页"] = homepage

            if metrics_trusted:
                row_data["S2匹配姓名"] = metrics.get("s2_name", "")
                row_data["S2作者ID"] = metrics.get("s2_author_id", "")
                row_data["S2主页"] = metrics.get("s2_url", "")
                row_data["S2代表论文"] = metrics.get("top_papers", "")
                _append_information_source(row_data, metrics.get("s2_url"))
            else:
                row_data["S2匹配姓名"] = ""
                row_data["S2作者ID"] = ""
                row_data["S2主页"] = ""
                row_data["S2代表论文"] = ""
        else:
            row_data["S2匹配姓名"] = ""
            row_data["S2作者ID"] = ""
            row_data["S2主页"] = ""
            row_data["S2代表论文"] = ""

        enriched_rows.append(row_data)

    return pd.DataFrame(enriched_rows)


def filter_deceased_experts(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df

    work_df = df.copy()
    if "独立生存状态核验" not in work_df.columns:
        work_df["独立生存状态核验"] = "未执行"
    deceased = work_df["独立生存状态核验"].fillna("").astype(str).str.strip().eq("确认已故")
    removed_count = int(deceased.sum())
    if removed_count:
        safe_print(f"[生存状态过滤] 已剔除 {removed_count} 位明确已故专家。")
    return work_df.loc[~deceased].copy()


def _domain_relevance_profile(query: str) -> tuple[str, dict] | tuple[None, None]:
    main_domain, _ = query_domain_parts(query)
    query_text = f"{main_domain} {query}".lower()
    for profile_name, profile in DOMAIN_RELEVANCE_PROFILES.items():
        if any(alias.lower() in query_text for alias in profile["aliases"]):
            return profile_name, profile
    return None, None


def filter_obvious_domain_mismatches(df: pd.DataFrame, query: str = "") -> pd.DataFrame:
    """
    保守剔除明显不属于用户目标大领域的专家。

    只有核心专家证据完全没有目标领域正向线索，同时命中一个强冲突领域词组时才剔除。
    不读取系统强制写入的“学科领域/细分领域”，避免这些标签掩盖真实错配。
    """
    if df.empty or not query:
        return df

    profile_name, profile = _domain_relevance_profile(query)
    if not profile:
        return df

    evidence_columns = [
        "研究兴趣",
        "工作单位",
        "职位",
        "工作经历",
        "教育背景",
        "主要成果",
        "入选依据",
        "OpenAlex主题",
        "S2代表论文",
    ]
    removed = []

    def should_keep(row: pd.Series) -> bool:
        evidence_text = " ".join(
            str(row.get(column, "") or "")
            for column in evidence_columns
        ).lower()
        if any(keyword.lower() in evidence_text for keyword in profile["positive"]):
            return True

        conflict_hits = []
        for conflict_group in profile["conflicts"]:
            hits = [keyword for keyword in conflict_group if keyword.lower() in evidence_text]
            if hits:
                conflict_hits.extend(hits)
        if not conflict_hits:
            return True

        removed.append(
            (
                str(row.get("专家姓名", row.get("姓名", "")) or "未知专家").strip(),
                "、".join(dict.fromkeys(conflict_hits[:4])),
            )
        )
        return False

    keep_mask = df.apply(should_keep, axis=1)
    removed_count = int((~keep_mask).sum())
    if removed_count:
        examples = "；".join(f"{name}({reason})" for name, reason in removed[:5])
        safe_print(
            f"[领域错配清洗] 目标大领域={profile_name}，已剔除 {removed_count} 条明显跨领域记录；"
            f"示例={examples}。"
        )
    return df.loc[keep_mask].copy()


def add_ranking_metrics(df: pd.DataFrame, query: str = "", include_chinese_experts: bool = False) -> pd.DataFrame:
    """
    应用用户设计的评价公式：
    TotalScore = H + (log10(Citations) * 15) + (i10 * 0.2) + TitleBonus
    """
    df = remove_non_expert_rows(df)
    df = df.drop(columns=[column for column in GENERATED_METRIC_COLUMNS if column in df.columns], errors="ignore")
    df = clean_contact_fields(df)
    df = filter_obvious_domain_mismatches(df, query)
    fast_demo_mode = os.environ.get(
        "EXPERTSEARCH_DEMO_FAST_MODE", ""
    ).strip().lower() in {"1", "true", "yes", "on"}
    if fast_demo_mode:
        safe_print(
            "[演示质量链路] 三分钟平衡模式：保留逐人 OpenAlex 指标，"
            "限量执行主页与合并 Tavily 补查；跳过高耗时深挖。"
        )
        stage_started = time.perf_counter()
        df = enrich_openalex_metrics(df, query)
        safe_print(
            f"[演示质量耗时] OpenAlex指标={time.perf_counter() - stage_started:.1f} 秒。"
        )
        stage_started = time.perf_counter()
        df = enrich_expert_details(df, query)
        safe_print(
            f"[演示质量耗时] 限量网页补全={time.perf_counter() - stage_started:.1f} 秒。"
        )
        df = normalize_homepage_access_status(df)
        df = clean_contact_fields(df)
        stage_started = time.perf_counter()
        df = verify_survival_status(df, query)
        safe_print(
            f"[演示质量耗时] 生存核验={time.perf_counter() - stage_started:.1f} 秒。"
        )
        df = filter_deceased_experts(df)
    else:
        df = enrich_openalex_metrics(df, query)
        df = enrich_semantic_scholar_metrics(df, query)
        df = enrich_expert_details(df, query)
        df = normalize_homepage_access_status(df)
        df = clean_contact_fields(df)
        df = verify_survival_status(df, query)
        df = filter_deceased_experts(df)
    if not include_chinese_experts:
        df = filter_foreign_experts(df)

    if df.empty:
        for column in RANKING_COLUMNS:
            df[column] = []
        return df

    ranking_rows = []
    for _, row in df.iterrows():
        h_index = parse_metric_number(row.get("H指数", 0))
        citations = estimate_citations(row)
        citation_score = math.log10(citations) * 15 if citations > 0 else 0.0
        i10_index = get_i10_metric(row)
        i10_score = i10_index * 0.2
        bonus, bonus_hits = title_bonus(row)
        total_score = h_index + citation_score + i10_score + bonus
        ranking_rows.append(
            {
                "评价_H指数": round(h_index, 2),
                "评价_总被引次数": int(citations) if citations else 0,
                "评价_顶级头衔标识": bonus_hits,
                "评价_i10指数": round(i10_index, 2),
                "评价_TotalScore": round(total_score, 2),
            }
        )

    ranking_df = pd.DataFrame(ranking_rows)
    df = pd.concat([df.reset_index(drop=True), ranking_df], axis=1)
    df = df.sort_values("评价_TotalScore", ascending=False, na_position="last")
    df = drop_near_duplicate_experts(df)
    df = add_name_verification_columns(df)
    return df.drop(columns=[column for column in INTERNAL_ENRICHMENT_COLUMNS if column in df.columns], errors="ignore")


def _has_severe_information_gaps(row: pd.Series) -> bool:
    identity_columns = ["个人主页", "工作单位", "职位"]
    detail_columns = [
        "邮箱/电话",
        "工作经历",
        "教育背景",
        "国内合作学者与单位",
        "入选依据",
        "主要成果",
    ]
    identity_missing = sum(_is_missing_or_pending(row.get(column, "")) for column in identity_columns)
    detail_missing = sum(_is_missing_or_pending(row.get(column, "")) for column in detail_columns)
    no_traceable_source = _row_source_urls(row) == "暂无可追溯网页来源"
    verification_pending = str(row.get("专家姓名验证", "")).strip() != "已验证"

    return (
        identity_missing >= 2
        or detail_missing >= 2
        or no_traceable_source
        or (verification_pending and identity_missing >= 1)
    )


FINAL_ENRICHMENT_FIELDS = [
    "个人主页",
    "邮箱/电话",
    "研究兴趣",
    "工作单位",
    "职位",
    "工作经历",
    "教育背景",
    "主要成果",
    "国内合作学者与单位",
    "入选依据",
    "H指数",
    "i10指数",
    "总被引次数",
]


def _final_missing_fields(row: pd.Series) -> list[str]:
    missing = []
    for column in FINAL_ENRICHMENT_FIELDS:
        if column in {"H指数", "i10指数", "总被引次数"}:
            is_missing = parse_metric_number(row.get(column, 0)) <= 0
        elif column == "个人主页":
            value = str(row.get(column, "") or "").strip()
            is_missing = not value.startswith(("http://", "https://"))
        elif column == "邮箱/电话":
            is_missing = not bool(re.search(r"[\w.+-]+@[\w.-]+\.[A-Za-z]{2,}", str(row.get(column, ""))))
        else:
            is_missing = _is_missing_or_pending(row.get(column, ""))
        if is_missing:
            missing.append(column)
    return missing


def _refresh_evaluation_columns(df: pd.DataFrame) -> pd.DataFrame:
    """外部补全更新指标后，重新计算评分列，避免评价列继续显示旧的暂无值。"""
    work_df = df.copy()
    for index, row in work_df.iterrows():
        h_index = parse_metric_number(row.get("H指数", 0))
        citations = estimate_citations(row)
        i10_index = get_i10_metric(row)
        bonus, bonus_hits = title_bonus(row)
        total_score = (
            h_index
            + (math.log10(citations) * 15 if citations > 0 else 0.0)
            + (i10_index * 0.2)
            + bonus
        )
        work_df.at[index, "评价_H指数"] = round(h_index, 2)
        work_df.at[index, "评价_总被引次数"] = int(citations) if citations else 0
        work_df.at[index, "评价_顶级头衔标识"] = bonus_hits
        work_df.at[index, "评价_i10指数"] = round(i10_index, 2)
        work_df.at[index, "评价_TotalScore"] = round(total_score, 2)
    return work_df


def normalize_homepage_access_status(df: pd.DataFrame) -> pd.DataFrame:
    """根据最终访问方式和失败原因统一主页状态，修复批次合并后的状态冲突。"""
    if df.empty:
        return df

    work_df = df.copy()
    for column, default in [
        ("成功访问主页", "未尝试"),
        ("主页访问方式", "未尝试"),
        ("主页访问失败原因", ""),
    ]:
        if column not in work_df.columns:
            work_df[column] = default

    for index, row in work_df.iterrows():
        homepage = str(row.get("个人主页", "") or "").strip()
        method = str(row.get("主页访问方式", "") or "").strip()
        failure = str(row.get("主页访问失败原因", "") or "").strip()
        valid_homepage = homepage.startswith(("http://", "https://"))
        has_failure = not _is_missing_or_pending(failure)

        if method in {"HTTP", "OpenCLI"} and not has_failure:
            status = "是"
        elif "失败" in method or has_failure:
            status = "否"
        elif not valid_homepage:
            status = "未尝试"
            method = "未尝试"
        else:
            status = str(row.get("成功访问主页", "") or "").strip() or "未尝试"

        work_df.at[index, "成功访问主页"] = status
        work_df.at[index, "主页访问方式"] = method or "未尝试"
    return work_df


def selectively_enrich_final_experts(df: pd.DataFrame, query: str = "") -> pd.DataFrame:
    """
    最终合并后，对任何关键交付字段仍有空白的专家进行定向补全。

    各批次可能为同一专家提供互补信息；完成字段级合并后，再找出身份信息、
    教育背景、合作证据、成果或来源严重缺失的记录，并逐条尝试外部补全。
    不设置补全人数上限；工具会按每位专家实际缺失字段分流，仍坚持无公开证据时不编造。
    """
    if df.empty:
        return df

    work_df = df.copy()
    if os.environ.get("EXPERTSEARCH_DEMO_FAST_MODE", "").strip().lower() in {
        "1", "true", "yes", "on",
    }:
        safe_print(
            "[最终定向补全] 快速演示模式跳过最终全量深挖，"
            "保留批次阶段已获得的 OpenAlex、网页和主页信息。"
        )
        return work_df

    before_missing = {
        column: int(work_df.apply(lambda row: column in _final_missing_fields(row), axis=1).sum())
        for column in FINAL_ENRICHMENT_FIELDS
    }
    selected_indices = work_df.index[work_df.apply(lambda row: bool(_final_missing_fields(row)), axis=1)]
    if len(selected_indices) == 0:
        safe_print("[最终定向补全] 所有关键交付字段均已有可用信息，无需补全。")
        return work_df

    selected_df = work_df.loc[selected_indices].copy()
    selected_df = enrich_expert_details(selected_df, query, exhaustive=True)
    selected_df.index = selected_indices

    for index in selected_indices:
        for column in selected_df.columns:
            if column not in work_df.columns:
                work_df[column] = ""
            work_df.at[index, column] = selected_df.at[index, column]

    after_missing = {
        column: int(work_df.apply(lambda row: column in _final_missing_fields(row), axis=1).sum())
        for column in FINAL_ENRICHMENT_FIELDS
    }
    improved = {
        column: before_missing[column] - after_missing[column]
        for column in FINAL_ENRICHMENT_FIELDS
        if before_missing[column] != after_missing[column]
    }
    safe_print(
        f"[最终定向补全] 已对 {len(selected_indices)} 位存在关键字段空白的专家执行补全；"
        f"字段改善={improved or '未找到新的可验证公开信息'}；"
        f"补全后仍缺={after_missing}。"
    )
    return _refresh_evaluation_columns(work_df)


def filter_low_confidence_experts(df: pd.DataFrame, query: str = "") -> pd.DataFrame:
    """
    只剔除具有明确错配或明确领域不相关证据的候选。

    缺少个人级来源、主页暂不可访问或姓名仍待复核的专家继续保留，
    由最终交付表中的验证状态和来源字段明确提示，避免过早删除可继续补全的有效专家。
    """
    df = filter_obvious_domain_mismatches(df, query)
    if df.empty:
        return df

    relevance_failure = re.compile(r"直接关联证据不足|领域关联不足|与.+关联证据不足")

    def should_keep(row: pd.Series) -> bool:
        verification = str(row.get("姓名验证依据", ""))
        if re.search(r"疑似错配|姓名不一致|身份错配", verification):
            return False

        row_text = " ".join(
            str(row.get(column, ""))
            for column in ["工作单位", "职位", "领域关联依据", "主要成果"]
        )

        if relevance_failure.search(row_text):
            return False
        return True

    keep_mask = df.apply(should_keep, axis=1)
    removed_count = int((~keep_mask).sum())
    if removed_count:
        safe_print(f"[最终质量门槛] 已剔除 {removed_count} 条明确身份错配或领域不相关的候选记录。")
    return df.loc[keep_mask].copy()


def finalize_merged_experts(
    df: pd.DataFrame,
    include_chinese_experts: bool = False,
    query: str = "",
) -> pd.DataFrame:
    """
    离线整理已经完成增强和评分的批次结果。

    先字段级合并同名专家的互补信息，再对所有信息严重缺失记录进行定向补全。
    最终写出前补查尚未完成独立生存状态核验的专家，并剔除确认已故者。
    """
    work_df = remove_non_expert_rows(df)
    work_df = clean_contact_fields(work_df)
    work_df = filter_deceased_experts(work_df)
    work_df = filter_obvious_domain_mismatches(work_df, query)
    if not include_chinese_experts:
        work_df = filter_foreign_experts(work_df)
    work_df = drop_near_duplicate_experts(work_df)
    work_df = selectively_enrich_final_experts(work_df, query)
    work_df = normalize_homepage_access_status(work_df)
    fast_demo_mode = os.environ.get(
        "EXPERTSEARCH_DEMO_FAST_MODE", ""
    ).strip().lower() in {"1", "true", "yes", "on"}
    work_df = verify_survival_status(
        work_df,
        query,
        exhaustive=not fast_demo_mode,
    )
    work_df = filter_deceased_experts(work_df)
    work_df = work_df.drop(columns=["专家姓名验证", "姓名验证依据"], errors="ignore")
    work_df = add_name_verification_columns(work_df)
    work_df = filter_low_confidence_experts(work_df, query)
    if "评价_TotalScore" in work_df.columns:
        work_df = work_df.sort_values("评价_TotalScore", ascending=False, na_position="last")
    return work_df.reset_index(drop=True)


def extract_markdown_table(text: str) -> list:
    """
    从文本中提取 Markdown 表格，返回一个包含行数据的列表（每一行是一个列表）
    """
    lines = text.split('\n')
    table_lines = []
    in_table = False
    
    for line in lines:
        line = line.strip()
        if line.startswith('|') and line.endswith('|'):
            in_table = True
            table_lines.append(line)
        elif in_table:
            # 如果脱离了表格边界，停止解析
            break
            
    if not table_lines:
        return []
        
    # 解析行数据
    parsed_table = []
    for row in table_lines:
        # 去掉首尾的 '|'
        row_content = row.strip('|')
        # 分割并去除空格
        cells = [cell.strip() for cell in row_content.split('|')]
        parsed_table.append(cells)
        
    # 去除分隔行 (通常是包含 '---' 的那一行)
    if len(parsed_table) > 1 and all(all(c == '-' or c == ':' for c in cell.replace(' ', '')) for cell in parsed_table[1] if cell):
        parsed_table.pop(1)
        
    return parsed_table

def excel_converter_node(state: AgentState):
    """
    负责数据清洗的节点，将智能体最终输出的信息转化为 Excel 表格
    """
    final_data = require_state_value(state, "final_data")
    query = state.get("query", "")
    include_chinese_experts = bool(state.get("include_chinese_experts", False))
    extra_data_source_urls = state.get("extra_data_source_urls", [])
    supplemental_document_names = state.get("supplemental_document_names", [])
    
    safe_print("\n" + "="*20 + " [调试：大模型原始输出] " + "="*20)
    safe_print(final_data)
    safe_print("="*64 + "\n")
    
    # 强化清洗1：无情剥离大模型可能附带的 markdown 代码块包裹标记
    final_data = final_data.replace("```markdown", "").replace("```", "").strip()
    
    # 强化清洗2：如果文本开头包含非表格的废话（如“好的，这是你要的表格：”），直接截取从第一个 '|' 开始的核心内容
    if '|' in final_data:
        final_data = final_data[final_data.find('|'):]
    
    # 批次文件由调用方写入临时目录；只有最终总表保留在 Expert_Results。
    output_dir = str(state.get("output_dir", "Expert_Results") or "Expert_Results")
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    # 2. 生成带时间戳的动态文件名
    timestamp = time.strftime("%Y%m%d_%H%M%S")
    file_name = f"Expert_Data_{timestamp}.xlsx"
    
    # 3. 将文件名拼接进目标文件夹路径中
    full_save_path = os.path.join(output_dir, file_name)
    
    from io import StringIO
    
    # 使用正则表达式精准捕捉以 | 开头并以 | 结尾的连续表格块
    # 这个正则可以无视表格前后的任何废话、说明或者 markdown 标记
    table_match = re.search(r'(\|.*\|(?:\n\|.*\|)+)', final_data)
    
    if table_match:
        clean_table_str = table_match.group(1)
        try:
            # 利用 StringIO 和 pandas 直接读取合法的 Markdown 表格字符串
            # 过滤掉 markdown 表格特有的分隔行 (如 |---|---|)
            lines = [line for line in clean_table_str.split('\n') if not re.match(r'^\|[-\s|:]+\|$', line.strip())]
            clean_table_str = '\n'.join(lines)
            
            df = pd.read_csv(StringIO(clean_table_str), sep='|', skipinitialspace=True).dropna(axis=1, how='all')
            # 去除列名和数据中可能存在的前后空格
            df.columns = df.columns.str.strip()
            for col in df.columns:
                if df[col].dtype == 'object':
                    df[col] = df[col].str.strip()

            main_domain, sub_domain = query_domain_parts(query)
            if "学科领域" not in df.columns:
                df["学科领域"] = main_domain
            # 每个批次由前端针对一个明确细分领域发起。这里强制写入该目标，
            # 避免模型自行输出宽泛领域或梯队条件，导致最终表无法区分归属。
            df["细分领域"] = normalize_subdomain_membership(sub_domain)

            df = add_ranking_metrics(df, query, include_chinese_experts=include_chinese_experts)
            safe_print(f"[系统提示] 成功提取出 {len(df)} 条专家数据，准备导出 Excel。")
            write_expert_excel(
                df,
                full_save_path,
                data_source_summary=data_source_summary(
                    extra_data_source_urls,
                    query=query,
                    supplemental_document_names=supplemental_document_names,
                ),
                query=query,
                extra_data_source_urls=extra_data_source_urls,
                supplemental_document_names=supplemental_document_names,
            )
            
        except Exception as e:
            safe_print(f"提取出表格文本但 Pandas 解析失败: {e}")
            write_empty_expert_table(full_save_path)
    else:
        safe_print("未能在最终数据中提取到有效的 Markdown 表格。")
        write_empty_expert_table(full_save_path)
        
    return {"excel_path": full_save_path}

