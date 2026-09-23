# -*- coding: utf-8 -*-

import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from expertsearch.table_completion import (
    _add_targeted_completion_evidence,
    _add_derived_completion_fields,
    _direct_updates,
    _extract_urls,
    _gate_validation_results,
    _row_sources,
    _run_completion_chunks,
    _run_completion_llm_batches,
    _sanitize_completion_evidence,
    _search_institution_hint,
    build_canonical_rows,
    completion_corrector_node,
    inspect_workbook_layout,
    write_completed_workbook,
)


def sample_workbook_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "推荐表"
    worksheet.merge_cells("A2:O2")
    worksheet["A2"] = "脑机接口领域全球顶尖人才推荐表"
    headers = [
        "序号",
        "姓名\n（中英文）",
        "性别",
        "所在地",
        "是否华人",
        "单位\n（中英文）",
        "职务",
        "主要教育经历",
        "主要工作经历",
        "细分领域",
        "推荐理由",
        "是否推荐为顶尖",
        "类型",
        "是否为\n顶尖人才",
        "备注",
    ]
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(4, column, header)
        cell.fill = PatternFill("solid", fgColor="2F6B4F")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    worksheet.append([])
    worksheet.cell(5, 1, 1)
    worksheet.cell(5, 2, "贺斌\nBin He")
    worksheet.cell(5, 6, "美国卡内基梅隆大学")
    worksheet.cell(6, 1, 2)
    worksheet.cell(6, 2, "测试专家\nTest Expert")
    worksheet.cell(6, 6, "Existing University")
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


class TableCompletionTests(unittest.TestCase):
    @patch("expertsearch.table_completion.table_completion_graph.invoke")
    def test_large_completion_is_split_into_resilient_chunks(self, mock_invoke):
        def fake_invoke(state):
            rows = state["source_rows"]
            return {
                "evidence_rows": [dict(row) for row in rows],
                "research_updates": [],
                "validation_results": [],
                "final_updates": [],
            }

        mock_invoke.side_effect = fake_invoke
        source_rows = [
            {"__excel_row__": row_number, "__source_name__": f"Expert {row_number}"}
            for row_number in range(5, 10)
        ]
        canonical_rows = [dict(row) for row in source_rows]
        with patch.dict(
            "os.environ",
            {
                "TABLE_COMPLETION_PROCESS_CHUNK_SIZE": "2",
                "TABLE_COMPLETION_PROCESS_CHUNK_RETRY_ATTEMPTS": "1",
            },
        ):
            result = _run_completion_chunks(
                "脑机接口",
                ["姓名"],
                source_rows,
                canonical_rows,
            )
        self.assertEqual(mock_invoke.call_count, 3)
        self.assertEqual(len(result["evidence_rows"]), 5)

    @patch("expertsearch.table_completion.table_completion_graph.invoke")
    def test_157_experts_are_processed_in_eleven_rounds_of_fifteen(self, mock_invoke):
        def fake_invoke(state):
            rows = state["source_rows"]
            return {
                "evidence_rows": [dict(row) for row in rows],
                "research_updates": [],
                "validation_results": [],
                "final_updates": [],
            }

        mock_invoke.side_effect = fake_invoke
        source_rows = [
            {"__excel_row__": row_number, "__source_name__": f"Expert {row_number}"}
            for row_number in range(5, 162)
        ]
        with patch.dict(
            "os.environ",
            {
                "TABLE_COMPLETION_PROCESS_CHUNK_SIZE": "15",
                "TABLE_COMPLETION_PROCESS_CHUNK_RETRY_ATTEMPTS": "1",
            },
        ):
            result = _run_completion_chunks(
                "脑机接口",
                ["姓名"],
                source_rows,
                [dict(row) for row in source_rows],
            )
        self.assertEqual(mock_invoke.call_count, 11)
        self.assertEqual(len(result["evidence_rows"]), 157)
        self.assertEqual(len(mock_invoke.call_args_list[-1].args[0]["source_rows"]), 7)

    @patch("expertsearch.table_completion.table_completion_graph.invoke")
    def test_round_retries_when_evidence_does_not_cover_every_expert(self, mock_invoke):
        source_rows = [
            {"__excel_row__": row_number, "__source_name__": f"Expert {row_number}"}
            for row_number in range(5, 8)
        ]
        mock_invoke.side_effect = [
            {
                "evidence_rows": [dict(source_rows[0])],
                "research_updates": [],
                "validation_results": [],
                "final_updates": [],
            },
            {
                "evidence_rows": [dict(row) for row in source_rows],
                "research_updates": [],
                "validation_results": [],
                "final_updates": [],
            },
        ]
        with patch.dict(
            "os.environ",
            {
                "TABLE_COMPLETION_PROCESS_CHUNK_SIZE": "15",
                "TABLE_COMPLETION_PROCESS_CHUNK_RETRY_ATTEMPTS": "2",
            },
        ):
            result = _run_completion_chunks(
                "脑机接口",
                ["姓名"],
                source_rows,
                [dict(row) for row in source_rows],
            )
        self.assertEqual(mock_invoke.call_count, 2)
        self.assertEqual(len(result["evidence_rows"]), 3)

    @patch("expertsearch.table_completion.invoke_llm_with_stage")
    def test_completion_llm_batch_retries_transient_connection_error(self, mock_invoke):
        mock_invoke.side_effect = [
            ConnectionError("temporary connection error"),
            SimpleNamespace(content='[{"row_number": 5, "updates": {"职务": "教授"}}]'),
        ]
        with patch.dict(
            "os.environ",
            {
                "TABLE_COMPLETION_LLM_RETRY_ATTEMPTS": "2",
                "TABLE_COMPLETION_LLM_RETRY_DELAY_SECONDS": "0",
            },
        ):
            result = _run_completion_llm_batches(
                "表格补全测试",
                "只输出 JSON",
                [{"row_number": 5}],
                1,
            )
        self.assertEqual(result[0]["updates"]["职务"], "教授")
        self.assertEqual(mock_invoke.call_count, 2)

    def test_detects_non_first_header_and_builds_canonical_rows(self):
        layout = inspect_workbook_layout(sample_workbook_bytes())
        self.assertEqual(layout.sheet_name, "推荐表")
        self.assertEqual(layout.header_row, 4)
        self.assertEqual(len(layout.data_rows), 2)

        canonical = build_canonical_rows(layout)
        self.assertEqual(canonical[0]["专家姓名"], "Bin He")
        self.assertEqual(canonical[0]["工作单位"], "美国卡内基梅隆大学")
        self.assertEqual(canonical[0]["__excel_row__"], 5)

    def test_writes_only_blank_cells_and_appends_audit_columns(self):
        source = sample_workbook_bytes()
        layout = inspect_workbook_layout(source)
        evidence_rows = [
            {
                "__excel_row__": 5,
                "信息来源": "https://example.edu/bin-he",
                "OpenAlex作者ID": "https://openalex.org/A123",
                "专家姓名验证": "已验证",
                "独立生存状态核验": "在世",
            },
            {
                "__excel_row__": 6,
                "信息来源": "https://example.edu/test",
                "专家姓名验证": "待复核",
                "独立生存状态核验": "待核验",
            },
        ]
        updates = [
            {
                "row_number": 5,
                "final_updates": {
                    "单位\n（中英文）": "不应覆盖原单位",
                    "职务": "教授",
                    "主要工作经历": "现任卡内基梅隆大学教授",
                },
            }
        ]
        with tempfile.TemporaryDirectory() as directory:
            output = str(Path(directory) / "completed.xlsx")
            validations = [
                {
                    "row_number": 5,
                    "accepted_updates": {
                        "职务": "教授",
                        "主要工作经历": "现任卡内基梅隆大学教授",
                    },
                    "field_sources": {
                        "职务": ["https://example.edu/bin-he"],
                        "主要工作经历": ["https://example.edu/bin-he"],
                    },
                }
            ]
            summary = write_completed_workbook(
                source, layout, updates, evidence_rows, output, validations
            )
            workbook = load_workbook(output)
            worksheet = workbook["推荐表"]

        self.assertEqual(worksheet["F5"].value, "美国卡内基梅隆大学")
        self.assertEqual(worksheet["G5"].value, "教授")
        self.assertEqual(worksheet["I5"].value, "现任卡内基梅隆大学教授")
        self.assertEqual(worksheet["P4"].value, "补全信息来源")
        self.assertIn("openalex.org", worksheet["P5"].value)
        self.assertIn("字段级验证通过 2 项", worksheet["Q5"].value)
        self.assertIn("补全证据明细", workbook.sheetnames)
        detail = workbook["补全证据明细"]
        self.assertEqual(detail["C2"].value, "职务")
        self.assertIn("example.edu", detail["F2"].value)
        self.assertEqual(summary["total_rows"], 2)
        self.assertEqual(summary["completed_rows"], 1)
        self.assertEqual(summary["filled_cells"], 2)

    @patch("expertsearch.table_completion.invoke_llm_with_stage")
    def test_corrector_cannot_restore_validator_rejected_columns(self, mock_invoke):
        mock_invoke.return_value = SimpleNamespace(
            content=(
                '[{"row_number": 5, "final_updates": '
                '{"职务": "教授", "性别": "男"}}]'
            )
        )
        state = {
            "query": "脑机接口",
            "columns": ["姓名\n（中英文）", "性别", "职务"],
            "source_rows": [
                {
                    "__excel_row__": 5,
                    "__source_name__": "贺斌\nBin He",
                    "姓名\n（中英文）": "贺斌\nBin He",
                    "性别": None,
                    "职务": None,
                }
            ],
            "evidence_rows": [
                {
                    "__excel_row__": 5,
                    "专家姓名": "Bin He",
                    "职位": "教授",
                }
            ],
            "research_updates": [
                {"row_number": 5, "updates": {"职务": "教授", "性别": "男"}}
            ],
            "validation_results": [
                {
                    "row_number": 5,
                    "accepted_updates": {"职务": "教授"},
                    "rejected_columns": {"性别": "证据未明确"},
                }
            ],
        }
        result = completion_corrector_node(state)
        self.assertEqual(
            result["final_updates"],
            [{"row_number": 5, "final_updates": {"职务": "教授"}}],
        )

    @patch("expertsearch.table_completion.invoke_llm_with_stage")
    def test_corrector_keeps_validator_accepted_field_when_llm_omits_it(self, mock_invoke):
        mock_invoke.return_value = SimpleNamespace(
            content='[{"row_number": 5, "final_updates": {}}]'
        )
        state = {
            "query": "脑机接口",
            "columns": ["姓名\n（中英文）", "职务"],
            "source_rows": [
                {
                    "__excel_row__": 5,
                    "__source_name__": "贺斌\nBin He",
                    "姓名\n（中英文）": "贺斌\nBin He",
                    "职务": None,
                }
            ],
            "evidence_rows": [{"__excel_row__": 5, "专家姓名": "Bin He"}],
            "research_updates": [{"row_number": 5, "updates": {"职务": "教授"}}],
            "validation_results": [
                {
                    "row_number": 5,
                    "accepted_updates": {"职务": "教授"},
                    "field_sources": {"职务": ["https://example.edu/bin-he"]},
                }
            ],
        }
        result = completion_corrector_node(state)
        self.assertEqual(
            result["final_updates"],
            [{"row_number": 5, "final_updates": {"职务": "教授"}}],
        )

    def test_completion_survival_gate_downgrades_namesake_death_evidence(self):
        source_row = {
            "__source_name__": "Elon Musk",
            "单位\n（中英文）": "Tesla / SpaceX",
        }
        evidence = {
            "__excel_row__": 5,
            "专家姓名": "Elon Musk",
            "独立生存状态核验": "确认已故",
            "生存状态": "已故",
            "生存状态核验依据": (
                "Elon Musk died peacefully | https://example-funeral-home.com/obituary/elon-musk"
            ),
            "信息来源": "https://example-funeral-home.com/obituary/elon-musk",
        }
        sanitized = _sanitize_completion_evidence(evidence, source_row)
        self.assertEqual(sanitized["补全身份锁定"], "待复核")
        self.assertEqual(sanitized["独立生存状态核验"], "待核验")
        self.assertEqual(sanitized["生存状态"], "")

    def test_json_evidence_urls_do_not_keep_trailing_quotes(self):
        value = (
            '[{"url": "https://profiles.stanford.edu/test"}, '
            '{"url": "https://openalex.org/A123"}]'
        )
        self.assertEqual(
            _extract_urls(value),
            ["https://profiles.stanford.edu/test", "https://openalex.org/A123"],
        )

    def test_locked_narrative_evidence_becomes_direct_proposal(self):
        source_row = {
            "__source_name__": "Jaimie Henderson",
            "职务": None,
            "主要工作经历": None,
            "细分领域": None,
        }
        evidence = {
            "补全身份锁定": "已锁定",
            "职位": "斯坦福大学神经外科教授",
            "工作经历": "现任斯坦福大学神经修复实验室主任。",
            "研究兴趣": "脑机接口、神经假体与功能神经外科",
            "个人主页": "https://profiles.stanford.edu/jaimie-henderson",
            "OpenAlex作者ID": "https://openalex.org/A123",
            "信息来源": (
                "https://profiles.stanford.edu/jaimie-henderson | "
                "https://openalex.org/A123"
            ),
        }
        updates = _direct_updates(
            source_row,
            evidence,
            ["职务", "主要工作经历", "细分领域"],
        )
        self.assertEqual(updates["职务"], "斯坦福大学神经外科教授")
        self.assertEqual(updates["主要工作经历"], "现任斯坦福大学神经修复实验室主任。")
        self.assertEqual(updates["细分领域"], "脑机接口、神经假体与功能神经外科")
        self.assertIn("profiles.stanford.edu", evidence["补全字段来源提示"])

    def test_identity_lock_requires_affiliation_or_official_profile(self):
        source_row = {
            "__source_name__": "Jaimie Henderson",
            "单位\n（中英文）": "Stanford University 美国斯坦福大学",
        }
        evidence = {
            "专家姓名": "Jaimie Henderson",
            "OpenAlex匹配姓名": "Jaimie Henderson",
            "OpenAlex作者ID": "https://openalex.org/A1",
            "OpenAlex匹配机构": "Stanford University",
            "个人主页": "https://profiles.stanford.edu/jaimie-henderson",
            "成功访问主页": "是",
        }
        sanitized = _sanitize_completion_evidence(evidence, source_row)
        self.assertEqual(sanitized["补全身份锁定"], "已锁定")

    def test_targeted_search_filters_namesakes_and_irrelevant_pages(self):
        class FakeTavily:
            calls = 0

            def invoke(self, _payload):
                self.calls += 1
                return {
                    "results": [
                        {
                            "title": "Jaimie Henderson - Stanford Medicine",
                            "url": "https://profiles.stanford.edu/jaimie-henderson",
                            "content": "Jaimie Henderson directs functional neurosurgery at Stanford University.",
                        },
                        {
                            "title": "Another Henderson",
                            "url": "https://example.edu/other",
                            "content": "Unrelated professor at another university.",
                        },
                    ]
                }

        source_row = {
            "__source_name__": "Jaimie Henderson",
            "单位\n（中英文）": "Stanford University",
            "职务": None,
            "主要教育经历": None,
        }
        tavily = FakeTavily()
        enriched = _add_targeted_completion_evidence(
            {"专家姓名": "Jaimie Henderson"},
            source_row,
            "脑机接口",
            tavily,
        )
        self.assertIn("profiles.stanford.edu", enriched["补全定向证据"])
        self.assertNotIn("example.edu/other", enriched["补全定向证据"])
        self.assertEqual(tavily.calls, 1)

    def test_search_institution_hint_uses_short_primary_institution(self):
        source_row = {
            "单位\n（中英文）": (
                "Brown University 布朗大学、"
                "Wyss Center for Bio and Neuroengineering 怀斯中心"
            )
        }
        self.assertEqual(_search_institution_hint(source_row), "Brown University 布朗大学")

    def test_derived_reason_and_top_marker_use_only_accepted_fields(self):
        results = [
            {
                "row_number": 5,
                "accepted_updates": {
                    "职务": "某大学神经工程教授",
                    "细分领域": "脑机接口与神经工程",
                    "是否推荐为顶尖": "是",
                },
                "field_sources": {
                    "职务": ["https://example.edu/profile"],
                    "细分领域": ["https://example.edu/profile"],
                    "是否推荐为顶尖": ["https://example.edu/award"],
                },
            }
        ]
        enriched = _add_derived_completion_fields(
            results,
            [
                {
                    "__excel_row__": 5,
                    "推荐理由": None,
                    "是否推荐为顶尖": None,
                    "是否为\n顶尖人才": None,
                }
            ],
            [{"__excel_row__": 5, "补全身份锁定": "已锁定"}],
            ["职务", "细分领域", "推荐理由", "是否推荐为顶尖", "是否为\n顶尖人才"],
        )
        accepted = enriched[0]["accepted_updates"]
        self.assertIn("脑机接口与神经工程", accepted["推荐理由"])
        self.assertEqual(accepted["是否为\n顶尖人才"], "是")

    def test_authoritative_honor_derives_both_top_markers(self):
        url = "https://example.edu/news/gruber-prize"
        results = [
            {
                "row_number": 5,
                "accepted_updates": {"备注": "2025年获Gruber Neuroscience Prize。"},
                "field_sources": {"备注": [url]},
            }
        ]
        enriched = _add_derived_completion_fields(
            results,
            [
                {
                    "__excel_row__": 5,
                    "备注": None,
                    "是否推荐为顶尖": None,
                    "是否为\n顶尖人才": None,
                }
            ],
            [{"__excel_row__": 5, "补全身份锁定": "已锁定"}],
            ["备注", "是否推荐为顶尖", "是否为\n顶尖人才"],
        )
        accepted = enriched[0]["accepted_updates"]
        self.assertEqual(accepted["是否推荐为顶尖"], "是")
        self.assertEqual(accepted["是否为\n顶尖人才"], "是")

    def test_low_trust_sources_and_clipped_english_are_rejected(self):
        evidence = {
            "信息来源": (
                "https://en.wikipedia.org/wiki/Test | "
                "https://example.edu/profile/test | https://youtube.com/watch?v=1"
            )
        }
        self.assertEqual(_row_sources(evidence), "https://example.edu/profile/test")

        gated = _gate_validation_results(
            [
                {
                    "row_number": 5,
                    "accepted_updates": {
                        "主要工作经历": "He is a professor and served as director...",
                        "职务": "教授",
                    },
                    "field_sources": {
                        "主要工作经历": ["https://example.edu/profile/test"],
                        "职务": ["https://en.wikipedia.org/wiki/Test"],
                    },
                }
            ],
            [{"__excel_row__": 5, "主要工作经历": None, "职务": None}],
            ["主要工作经历", "职务"],
        )
        self.assertEqual(gated[0]["accepted_updates"], {})

    def test_pending_identity_rejects_commercial_biography_source(self):
        gated = _gate_validation_results(
            [
                {
                    "row_number": 5,
                    "accepted_updates": {"主要工作经历": "创办多家科技企业。"},
                    "field_sources": {
                        "主要工作经历": [
                            "https://www.abebooks.com/example-biography"
                        ]
                    },
                }
            ],
            [{"__excel_row__": 5, "主要工作经历": None}],
            ["主要工作经历"],
            [
                {
                    "__excel_row__": 5,
                    "信息来源": "https://www.abebooks.com/example-biography",
                    "补全身份锁定": "待复核",
                }
            ],
        )
        self.assertEqual(gated[0]["accepted_updates"], {})

    def test_pending_identity_allows_successful_same_domain_homepage(self):
        url = "https://www.neurotechcenter.org/people/test-expert"
        gated = _gate_validation_results(
            [
                {
                    "row_number": 5,
                    "accepted_updates": {"职务": "研究科学家"},
                    "field_sources": {"职务": [url]},
                }
            ],
            [{"__excel_row__": 5, "职务": None}],
            ["职务"],
            [
                {
                    "__excel_row__": 5,
                    "信息来源": url,
                    "个人主页": url,
                    "成功访问主页": "是",
                    "补全身份锁定": "待复核",
                }
            ],
        )
        self.assertEqual(gated[0]["accepted_updates"], {"职务": "研究科学家"})

    def test_generic_degree_without_school_is_rejected(self):
        gated = _gate_validation_results(
            [
                {
                    "row_number": 5,
                    "accepted_updates": {"主要教育经历": "博士学位"},
                    "field_sources": {
                        "主要教育经历": ["https://profiles.stanford.edu/test"]
                    },
                }
            ],
            [{"__excel_row__": 5, "主要教育经历": None}],
            ["主要教育经历"],
            [
                {
                    "__excel_row__": 5,
                    "信息来源": "https://profiles.stanford.edu/test",
                    "补全身份锁定": "已锁定",
                }
            ],
        )
        self.assertEqual(gated[0]["accepted_updates"], {})

    def test_name_plus_degree_without_school_is_rejected(self):
        gated = _gate_validation_results(
            [
                {
                    "row_number": 5,
                    "accepted_updates": {"主要教育经历": "Jaimie Henderson, MD"},
                    "field_sources": {
                        "主要教育经历": ["https://profiles.stanford.edu/test"]
                    },
                }
            ],
            [{"__excel_row__": 5, "主要教育经历": None}],
            ["主要教育经历"],
            [
                {
                    "__excel_row__": 5,
                    "信息来源": "https://profiles.stanford.edu/test",
                    "补全身份锁定": "已锁定",
                }
            ],
        )
        self.assertEqual(gated[0]["accepted_updates"], {})

    def test_pronoun_only_does_not_support_gender(self):
        url = "https://profiles.example.edu/edward"
        gated = _gate_validation_results(
            [
                {
                    "row_number": 5,
                    "accepted_updates": {"性别": "男"},
                    "field_sources": {"性别": [url]},
                }
            ],
            [{"__excel_row__": 5, "性别": None}],
            ["性别"],
            [
                {
                    "__excel_row__": 5,
                    "信息来源": url,
                    "补全身份锁定": "已锁定",
                    "补全定向证据": (
                        '[{"url": "https://profiles.example.edu/edward", '
                        '"title": "Edward Chang", '
                        '"snippet": "He is a professor of neurosurgery."}]'
                    ),
                }
            ],
        )
        self.assertEqual(gated[0]["accepted_updates"], {})

    def test_vague_english_honor_reason_is_rejected(self):
        gated = _gate_validation_results(
            [
                {
                    "row_number": 5,
                    "accepted_updates": {
                        "推荐理由": "因具有可验证的权威头衔入选：National Academy"
                    },
                    "field_sources": {
                        "推荐理由": ["https://example.edu/profile"]
                    },
                }
            ],
            [{"__excel_row__": 5, "推荐理由": None}],
            ["推荐理由"],
            [
                {
                    "__excel_row__": 5,
                    "信息来源": "https://example.edu/profile",
                    "补全身份锁定": "已锁定",
                }
            ],
        )
        self.assertEqual(gated[0]["accepted_updates"], {})


if __name__ == "__main__":
    unittest.main()
