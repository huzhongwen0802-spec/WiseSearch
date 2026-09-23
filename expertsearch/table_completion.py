# -*- coding: utf-8 -*-

from __future__ import annotations

import io
import json
import os
import re
import time
from copy import copy
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, NotRequired, TypedDict
from urllib.parse import urlparse

import pandas as pd
from dotenv import load_dotenv
from langchain_core.messages import HumanMessage, SystemMessage
from langchain_tavily import TavilySearch
from .tavily_client import (
    configure_tavily_budget,
    create_tavily_search,
    invoke_tavily_search,
    tavily_is_configured,
)
from langgraph.graph import END, StateGraph
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

from .agents import invoke_llm_with_stage
from .expert_enrichment import enrich_expert_details
from .openalex_client import get_openalex_author_metrics
from .safe_logging import safe_print
from .survival_verification import verify_survival_status
from .utils import (
    EXPERT_TABLE_HEADERS,
    add_name_verification_columns,
    clean_contact_fields,
    enrich_openalex_metrics,
    enrich_semantic_scholar_metrics,
    normalize_homepage_access_status,
)

load_dotenv()


PLACEHOLDER_TERMS = {
    "",
    "-",
    "--",
    "/",
    "暂无",
    "暂无公开信息",
    "未知",
    "不详",
    "待补充",
    "nan",
    "none",
    "null",
}

NAME_HEADER_ALIASES = {
    "姓名",
    "姓名中英文",
    "专家姓名",
    "人才姓名",
    "name",
    "expertname",
}

HEADER_ROLE_ALIASES = {
    "专家姓名": NAME_HEADER_ALIASES,
    "国籍": {"国籍", "国家", "nationality", "country"},
    "个人主页": {"个人主页", "主页", "个人网站", "homepage", "website"},
    "邮箱/电话": {"邮箱电话", "联系方式", "邮箱", "电子邮箱", "email", "contact"},
    "研究兴趣": {
        "研究兴趣", "研究方向", "学科领域", "领域", "细分领域", "researchinterests"
    },
    "工作单位": {"单位", "单位中英文", "工作单位", "所在工作机构", "所在单位", "机构", "affiliation"},
    "职位": {"职位", "职务", "职称", "position", "title"},
    "工作经历": {"工作经历", "主要工作经历", "职业经历", "employmenthistory"},
    "教育背景": {"教育背景", "教育经历", "主要教育经历", "education"},
    "H指数": {"h指数", "hindex"},
    "主要成果": {"主要成果", "代表成果", "科研成果", "achievements"},
    "国内合作学者与单位": {"国内合作学者与单位", "国内合作", "中国合作"},
    "入选依据": {"入选依据", "推荐理由", "推荐依据", "selectionbasis"},
    "领域关联依据": {"领域关联依据", "领域关联", "相关性依据"},
    "生存状态": {"生存状态", "在世状态"},
    "信息来源": {"信息来源", "数据来源", "来源", "source"},
}

DIRECT_FILL_PRIORITY = {
    "国籍": ["国籍"],
    "个人主页": ["个人主页"],
    "邮箱/电话": ["邮箱/电话"],
    "研究兴趣": ["研究兴趣", "OpenAlex主题"],
    "工作单位": ["工作单位"],
    "职位": ["职位"],
    "工作经历": ["工作经历"],
    "教育背景": ["教育背景"],
    "H指数": ["H指数"],
    "主要成果": ["主要成果"],
    "国内合作学者与单位": ["国内合作学者与单位"],
    "入选依据": ["入选依据", "主要成果"],
    "领域关联依据": ["领域关联依据", "研究兴趣"],
    "生存状态": ["独立生存状态核验", "生存状态"],
    "信息来源": ["信息来源"],
}

EVIDENCE_FIELDS = [
    "专家姓名",
    "国籍",
    "个人主页",
    "邮箱/电话",
    "研究兴趣",
    "工作单位",
    "职位",
    "工作经历",
    "教育背景",
    "H指数",
    "i10指数",
    "总被引次数",
    "主要成果",
    "国内合作学者与单位",
    "入选依据",
    "领域关联依据",
    "信息来源",
    "OpenAlex匹配姓名",
    "OpenAlex作者ID",
    "OpenAlex匹配机构",
    "OpenAlex主题",
    "S2匹配姓名",
    "S2作者ID",
    "S2主页",
    "S2代表论文",
    "成功访问主页",
    "主页访问方式",
    "专家姓名验证",
    "姓名验证依据",
    "独立生存状态核验",
    "生存状态核验依据",
    "补全定向证据",
    "补全字段来源提示",
]

AUDIT_COLUMNS = ["补全信息来源", "补全验证状态", "补全说明"]

LOW_TRUST_SOURCE_DOMAINS = {
    "wikipedia.org",
    "youtube.com",
    "youtu.be",
    "scribd.com",
    "radaris.com",
    "linkedin.com",
    "businessinsider.com",
    "leadiq.com",
    "wikihow.com",
    "movavi.io",
    "coursmos.com",
    "vedantu.com",
    "zoominfo.com",
    "rocketreach.co",
    "facebook.com",
    "instagram.com",
    "x.com",
    "twitter.com",
    "abebooks.com",
    "baike.baidu.com",
}

AUTHORITATIVE_SOURCE_HINTS = {
    "openalex.org",
    "semanticscholar.org",
    "orcid.org",
    "nih.gov",
    "ncbi.nlm.nih.gov",
    "ieee.org",
    "acm.org",
    "nationalacademies.org",
    "royalsociety.org",
}

NARRATIVE_ROLES = {
    "工作单位",
    "职位",
    "研究兴趣",
    "工作经历",
    "教育背景",
    "主要成果",
    "国内合作学者与单位",
    "入选依据",
    "领域关联依据",
}


@dataclass(frozen=True)
class WorkbookLayout:
    sheet_name: str
    header_row: int
    columns: list[tuple[int, str]]
    data_rows: list[dict[str, Any]]


class TableCompletionState(TypedDict):
    query: str
    columns: list[str]
    source_rows: list[dict[str, Any]]
    canonical_rows: list[dict[str, Any]]
    evidence_rows: NotRequired[list[dict[str, Any]]]
    research_updates: NotRequired[list[dict[str, Any]]]
    validation_results: NotRequired[list[dict[str, Any]]]
    final_updates: NotRequired[list[dict[str, Any]]]


def _normalize_header(value: Any) -> str:
    return re.sub(r"[^0-9a-zA-Z一-鿿]+", "", str(value or "")).lower()


def _is_placeholder(value: Any) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return True
    return str(value).strip().lower() in PLACEHOLDER_TERMS


def _header_role(header: str) -> str | None:
    normalized = _normalize_header(header)
    for role, aliases in HEADER_ROLE_ALIASES.items():
        if normalized in aliases:
            return role
    return None


def _preferred_search_name(value: Any) -> str:
    text = str(value or "").strip()
    if not text:
        return ""
    parts = [part.strip() for part in re.split(r"[\n/；;]+", text) if part.strip()]
    latin_parts = [
        part
        for part in parts
        if len(re.findall(r"[A-Za-z]+", part)) >= 2
    ]
    return latin_parts[-1] if latin_parts else parts[0]


def inspect_workbook_layout(workbook_bytes: bytes) -> WorkbookLayout:
    """识别最可能的数据工作表、非首行表头和专家数据行。"""
    workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False)
    best: tuple[float, Any, int, list[tuple[int, str]]] | None = None

    for worksheet in workbook.worksheets:
        scan_rows = min(30, worksheet.max_row)
        scan_columns = min(80, worksheet.max_column)
        for row_number in range(1, scan_rows + 1):
            columns = [
                (column_number, str(worksheet.cell(row_number, column_number).value).strip())
                for column_number in range(1, scan_columns + 1)
                if not _is_placeholder(worksheet.cell(row_number, column_number).value)
            ]
            if not columns:
                continue
            normalized_headers = {_normalize_header(header) for _, header in columns}
            has_name = bool(normalized_headers & NAME_HEADER_ALIASES)
            known_roles = sum(1 for _, header in columns if _header_role(header))
            score = (100 if has_name else 0) + known_roles * 8 + min(len(columns), 20)
            if best is None or score > best[0]:
                best = (score, worksheet, row_number, columns)

    if best is None or best[0] < 100:
        raise ValueError("未识别到包含“姓名/专家姓名”字段的有效表头。")

    _, worksheet, header_row, columns = best
    name_columns = [
        column_number
        for column_number, header in columns
        if _normalize_header(header) in NAME_HEADER_ALIASES
    ]
    if not name_columns:
        raise ValueError("表格缺少可识别的专家姓名列。")
    name_column = name_columns[0]

    data_rows: list[dict[str, Any]] = []
    for row_number in range(header_row + 1, worksheet.max_row + 1):
        name_value = worksheet.cell(row_number, name_column).value
        if _is_placeholder(name_value):
            continue
        row = {header: worksheet.cell(row_number, column_number).value for column_number, header in columns}
        row["__excel_row__"] = row_number
        row["__source_name__"] = str(name_value).strip()
        data_rows.append(row)

    if not data_rows:
        raise ValueError("已识别表头，但未发现包含专家姓名的数据行。")
    return WorkbookLayout(worksheet.title, header_row, columns, data_rows)


def build_canonical_rows(layout: WorkbookLayout) -> list[dict[str, Any]]:
    canonical_rows = []
    for source_row in layout.data_rows:
        row = {header: "" for header in EXPERT_TABLE_HEADERS}
        for header, value in source_row.items():
            if header.startswith("__") or _is_placeholder(value):
                continue
            role = _header_role(header)
            if role:
                row[role] = value
        row["专家姓名"] = _preferred_search_name(source_row.get("__source_name__"))
        row["__excel_row__"] = source_row["__excel_row__"]
        row["__source_name__"] = source_row["__source_name__"]
        canonical_rows.append(row)
    return canonical_rows


def _clip(value: Any, limit: int = 900) -> str:
    text = " ".join(str(value or "").split())
    return text if len(text) <= limit else text[:limit].rstrip() + "..."


def _normalize_person_name(value: Any) -> str:
    tokens = re.findall(r"[A-Za-z\u4e00-\u9fff]+", str(value or "").lower())
    return " ".join(token for token in tokens if len(token) > 1)


def _same_person_name(left: Any, right: Any) -> bool:
    left_name = _normalize_person_name(left)
    right_name = _normalize_person_name(right)
    if not left_name or not right_name:
        return False
    if left_name == right_name:
        return True
    left_tokens = set(left_name.split())
    right_tokens = set(right_name.split())
    return len(left_tokens & right_tokens) >= 2 and (
        left_tokens.issubset(right_tokens) or right_tokens.issubset(left_tokens)
    )


def _extract_urls(value: Any) -> list[str]:
    matches = re.findall(
        r"https?://[^\s|；;，,\"'<>\]\}]+",
        str(value or ""),
        flags=re.IGNORECASE,
    )
    return [url.rstrip(".。:：!?！？)）") for url in matches]


def _source_domain(url: str) -> str:
    try:
        return urlparse(url).netloc.lower().split(":", 1)[0].removeprefix("www.")
    except ValueError:
        return ""


def _domain_matches(domain: str, candidate: str) -> bool:
    return domain == candidate or domain.endswith(f".{candidate}")


def _source_tier(url: str) -> int:
    domain = _source_domain(url)
    if not domain:
        return 4
    if any(_domain_matches(domain, item) for item in LOW_TRUST_SOURCE_DOMAINS):
        return 4
    if (
        domain.endswith(".gov")
        or domain.endswith(".gov.cn")
        or domain.endswith(".edu")
        or ".edu." in domain
        or domain.endswith(".ac.cn")
        or domain.endswith(".ac.uk")
        or any(_domain_matches(domain, item) for item in AUTHORITATIVE_SOURCE_HINTS)
    ):
        return 1
    if domain.endswith(".org") or domain.endswith(".org.cn"):
        return 2
    return 3


def _ranked_sources(evidence_row: dict[str, Any], limit: int = 5) -> list[str]:
    urls: list[str] = []
    for field in [
        "信息来源", "OpenAlex作者ID", "S2主页", "个人主页", "生存状态核验依据", "补全定向证据"
    ]:
        urls.extend(_extract_urls(evidence_row.get(field)))
    unique_urls = list(dict.fromkeys(urls))
    trusted = [url for url in unique_urls if _source_tier(url) < 4]
    trusted.sort(key=lambda url: (_source_tier(url), unique_urls.index(url)))
    return trusted[:limit]


def _completion_core_sources(evidence_row: dict[str, Any], limit: int = 12) -> list[str]:
    urls: list[str] = []
    for field in [
        "OpenAlex作者ID",
        "S2主页",
        "个人主页",
        "补全定向证据",
        "生存状态核验依据",
    ]:
        urls.extend(_extract_urls(evidence_row.get(field)))
    unique = [url for url in dict.fromkeys(urls) if _source_tier(url) < 4]
    unique.sort(key=lambda url: (_source_tier(url), urls.index(url)))
    return unique[:limit]


def _institution_tokens(value: Any) -> set[str]:
    text = str(value or "").lower()
    stopwords = {
        "university", "college", "institute", "institution", "school", "department",
        "academy", "center", "centre", "laboratory", "lab", "the", "of", "and",
        "大学", "学院", "研究院", "研究所", "实验室", "中心", "公司", "集团",
    }
    return {
        token
        for token in re.findall(r"[a-z]{3,}|[\u4e00-\u9fff]{2,}", text)
        if token not in stopwords
    }


def _source_institution(source_row: dict[str, Any]) -> str:
    return " ".join(
        str(source_row.get(column, "") or "")
        for column in source_row
        if _header_role(column) == "工作单位"
    ).strip()


def _search_institution_hint(source_row: dict[str, Any]) -> str:
    """为搜索生成短机构提示，避免整段中英双语单位被精确匹配。"""
    institution = _source_institution(source_row)
    if not institution:
        return ""
    parts = [
        part.strip(" （）()[]")
        for part in re.split(r"[、，,；;|/\n]+", institution)
        if part.strip(" （）()[]")
    ]
    latin = [
        part for part in parts
        if len(re.findall(r"[A-Za-z]", part)) >= 5
        and any(
            token in part.lower()
            for token in ["university", "institute", "center", "centre", "college", "company"]
        )
    ]
    candidate = latin[0] if latin else (parts[0] if parts else institution)
    return _clip(candidate, 90)


def _institutions_match(source_institution: Any, matched_institution: Any) -> bool:
    source_tokens = _institution_tokens(source_institution)
    matched_tokens = _institution_tokens(matched_institution)
    if not source_tokens or not matched_tokens:
        return False
    shared = source_tokens & matched_tokens
    if len(shared) >= 2:
        return True
    return any(len(token) >= 5 for token in shared)


def _is_official_profile_url(url: str) -> bool:
    domain = _source_domain(url)
    if not domain or _source_tier(url) >= 4:
        return False
    return (
        domain.endswith(".edu")
        or ".edu." in domain
        or domain.endswith(".gov")
        or domain.endswith(".gov.cn")
        or domain.endswith(".ac.cn")
        or domain.endswith(".ac.uk")
    )


def _mentions_person(name: Any, text: Any) -> bool:
    normalized = _normalize_person_name(_preferred_search_name(name))
    haystack = str(text or "").lower()
    if not normalized:
        return False
    if re.search(r"[\u4e00-\u9fff]", normalized):
        return normalized.replace(" ", "") in re.sub(r"\s+", "", haystack)
    tokens = normalized.split()
    if len(tokens) == 1:
        return tokens[0] in haystack
    return tokens[0] in haystack and tokens[-1] in haystack


def _missing_search_groups(source_row: dict[str, Any]) -> list[tuple[str, str]]:
    missing_headers = [
        header for header, value in source_row.items()
        if not header.startswith("__") and _is_placeholder(value)
    ]
    normalized = " ".join(_normalize_header(header) for header in missing_headers)
    groups: list[tuple[str, str]] = []
    if any(token in normalized for token in ["性别", "所在地", "国籍", "是否华人"]):
        groups.append(("身份与所在地", "official biography nationality location gender"))
    if any(
        token in normalized
        for token in ["职务", "职位", "教育", "工作经历", "职业经历"]
    ):
        groups.append(("教育与任职", "official profile CV education degree career position"))
    if any(
        token in normalized
        for token in ["细分领域", "推荐理由", "顶尖", "类型", "备注", "成果", "奖项"]
    ):
        groups.append(("领域与荣誉", "research interests awards honors fellow academy achievements"))
    return groups


def _make_completion_tavily() -> TavilySearch | None:
    if not tavily_is_configured():
        return None
    try:
        return create_tavily_search(
            max_results=max(1, int(os.environ.get("TABLE_COMPLETION_TARGETED_SEARCH_RESULTS", "8"))),
            search_depth="advanced",
            topic="general",
            include_answer=False,
            handle_tool_error=True,
        )
    except Exception as exc:
        safe_print(f"[表格补全工具状态] Tavily 初始化失败，本轮使用其他工具继续: {exc}")
        return None


def _targeted_search_items(tavily: TavilySearch, query: str) -> list[dict[str, str]]:
    result = invoke_tavily_search(tavily, query, stage="已有表格信息补全")
    if not isinstance(result, dict):
        return []
    return [
        {
            "title": str(item.get("title") or ""),
            "url": str(item.get("url") or ""),
            "content": str(item.get("content") or item.get("raw_content") or ""),
        }
        for item in result.get("results", [])
    ]


def _add_targeted_completion_evidence(
    evidence_row: dict[str, Any],
    source_row: dict[str, Any],
    query: str,
    tavily: TavilySearch | None,
) -> dict[str, Any]:
    if tavily is None:
        return evidence_row
    groups = _missing_search_groups(source_row)
    max_queries = max(
        0,
        int(os.environ.get("TABLE_COMPLETION_TARGETED_SEARCH_MAX_QUERIES_PER_ROW", "3")),
    )
    if not groups or max_queries <= 0:
        return evidence_row

    compact_search = os.environ.get(
        "TABLE_COMPLETION_COMPACT_TARGETED_SEARCH", "true"
    ).lower() == "true"
    if compact_search:
        search_groups = [
            (
                "综合补全",
                "official profile biography CV education degree career position "
                "research interests awards honors fellow academy achievements",
            )
        ]
    else:
        search_groups = groups[:max_queries]

    name = _preferred_search_name(source_row.get("__source_name__", ""))
    institution = _source_institution(source_row)
    institution_hint = _search_institution_hint(source_row)
    institution_tokens = _institution_tokens(institution)
    query_tokens = _institution_tokens(query)
    homepage_domains = {
        _source_domain(url)
        for url in _extract_urls(evidence_row.get("个人主页", ""))
        if _source_domain(url)
    }
    accepted: list[dict[str, str]] = []
    for group_name, terms in search_groups:
        institution_clause = f' "{institution_hint}"' if institution_hint else ""
        search_query = f'"{name}"{institution_clause} {terms} {query}'.strip()
        try:
            items = _targeted_search_items(tavily, search_query)
        except Exception as exc:
            safe_print(f"[表格补全定向检索警告] {name} | {group_name} | {exc}")
            continue
        for item in items:
            url = item["url"]
            combined = f"{item['title']} {item['content']} {url}"
            if _source_tier(url) >= 4 or not _mentions_person(name, combined):
                continue
            lowered = combined.lower()
            context_match = any(token in lowered for token in institution_tokens | query_tokens)
            same_homepage_domain = _source_domain(url) in homepage_domains
            if not context_match and not _is_official_profile_url(url) and not same_homepage_domain:
                continue
            accepted.append(
                {
                    "group": group_name,
                    "title": _clip(item["title"], 180),
                    "url": url,
                    "snippet": _clip(item["content"], 520),
                    "tier": str(_source_tier(url)),
                }
            )

    unique: list[dict[str, str]] = []
    seen_urls: set[str] = set()
    for item in sorted(accepted, key=lambda value: int(value["tier"])):
        if item["url"] in seen_urls:
            continue
        seen_urls.add(item["url"])
        unique.append(item)
    if unique:
        evidence_row["补全定向证据"] = json.dumps(unique[:10], ensure_ascii=False)
        existing_sources = _extract_urls(evidence_row.get("信息来源", ""))
        evidence_row["信息来源"] = " | ".join(
            dict.fromkeys(existing_sources + [item["url"] for item in unique[:10]])
        )
    safe_print(
        f"[表格补全定向检索] {name} | 查询 {len(search_groups)} 组 | "
        f"保留相关证据 {len(unique[:10])} 条。"
    )
    return evidence_row


def _has_identity_context(source_row: dict[str, Any], evidence_text: str) -> bool:
    institution = _source_institution(source_row)
    tokens = _institution_tokens(institution)
    lowered = evidence_text.lower()
    return bool(tokens and any(token in lowered for token in tokens))


def _is_low_quality_completion_value(column: str, value: Any) -> bool:
    text = " ".join(str(value or "").split()).strip()
    if _is_placeholder(text):
        return True
    if re.search(
        r"(?:\.\.\.|…)$|error sending|connection error|operation timed out|search results?",
        text,
        flags=re.IGNORECASE,
    ):
        return True
    if "教育" in _normalize_header(column) and re.fullmatch(
        r"(?:博士|硕士|学士|医学博士|哲学博士|博士学位|硕士学位|学士学位|ph\.?d\.?|m\.?d\.?)",
        text,
        flags=re.IGNORECASE,
    ):
        return True
    if "教育" in _normalize_header(column) and not re.search(
        r"大学|学院|学校|研究所|医学院|医院|培训|博士后|"
        r"university|college|school|institute|academy|hospital|medical center|postdoc",
        text,
        flags=re.IGNORECASE,
    ):
        return True
    role = _header_role(column)
    if role == "入选依据":
        chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", text))
        if (
            chinese_chars < 18
            or re.search(
                r"(?:National Academy|Academy|Fellow|award|prize)\s*$",
                text,
                flags=re.IGNORECASE,
            )
        ):
            return True
    if role in NARRATIVE_ROLES:
        latin_letters = len(re.findall(r"[A-Za-z]", text))
        chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", text))
        if latin_letters >= 30 and chinese_chars < 6:
            return True
    return False


def _explicit_sensitive_value_supported(
    column: str,
    value: Any,
    evidence_row: dict[str, Any],
    source_urls: list[str],
) -> bool:
    """敏感属性必须由同一条定向证据明确陈述，不能根据姓名或代词推断。"""
    try:
        items = json.loads(str(evidence_row.get("补全定向证据", "") or "[]"))
    except json.JSONDecodeError:
        return False
    if not isinstance(items, list):
        return False
    normalized_column = _normalize_header(column)
    normalized_value = str(value or "").strip().lower()
    for item in items:
        if not isinstance(item, dict) or item.get("url") not in source_urls:
            continue
        text = f"{item.get('title', '')} {item.get('snippet', '')}".lower()
        if "性别" in normalized_column:
            if normalized_value in {"男", "男性", "male"} and re.search(
                r"性别\s*[:：]?\s*男|男性|\b(?:sex|gender)\s*[:：]?\s*male\b",
                text,
            ):
                return True
            if normalized_value in {"女", "女性", "female"} and re.search(
                r"性别\s*[:：]?\s*女|女性|\b(?:sex|gender)\s*[:：]?\s*female\b",
                text,
            ):
                return True
        if "是否华人" in normalized_column and normalized_value in {"是", "华人", "华裔"}:
            if re.search(r"华人|华裔|华侨|chinese[- ]american|chinese[- ]born", text):
                return True
    return False


def _sanitize_completion_evidence(
    evidence_row: dict[str, Any],
    source_row: dict[str, Any],
) -> dict[str, Any]:
    sanitized = dict(evidence_row)
    source_name = source_row.get("__source_name__", "")
    matched_names = [
        sanitized.get("OpenAlex匹配姓名", ""),
        sanitized.get("S2匹配姓名", ""),
    ]
    mismatches = [
        str(name)
        for name in matched_names
        if not _is_placeholder(name) and not _same_person_name(source_name, name)
    ]
    if mismatches:
        for field in [
            "H指数", "i10指数", "总被引次数", "OpenAlex匹配姓名", "OpenAlex作者ID",
            "OpenAlex匹配机构", "OpenAlex主题", "S2匹配姓名", "S2作者ID", "S2主页", "S2代表论文",
        ]:
            sanitized[field] = ""
        sanitized["补全身份锁定"] = "未通过"
        sanitized["补全身份锁定依据"] = f"数据库匹配姓名与原表不一致：{'、'.join(mismatches)}"
    else:
        database_signals = sum(
            1 for value in [sanitized.get("OpenAlex作者ID"), sanitized.get("S2作者ID")]
            if not _is_placeholder(value)
        )
        homepage_verified = str(sanitized.get("成功访问主页", "")).strip() == "是"
        affiliation_match = _institutions_match(
            _source_institution(source_row),
            sanitized.get("OpenAlex匹配机构", ""),
        )
        official_profile_urls = [
            url for url in _completion_core_sources(sanitized, limit=20)
            if _is_official_profile_url(url)
        ]
        identity_locked = database_signals >= 1 and (
            affiliation_match or (homepage_verified and bool(official_profile_urls))
        )
        sanitized["补全身份锁定"] = "已锁定" if identity_locked else "待复核"
        sanitized["补全身份锁定依据"] = (
            f"同名数据库信号 {database_signals} 个；"
            f"OpenAlex机构{'匹配' if affiliation_match else '未匹配'}；"
            f"机构官方主页证据 {len(official_profile_urls)} 个；"
            f"主页{'访问成功' if homepage_verified else '未完成有效访问'}"
        )

    ranked_sources = _completion_core_sources(sanitized, limit=8)
    sanitized["信息来源"] = " | ".join(ranked_sources)
    sanitized["补全来源等级"] = (
        "高可信来源优先" if any(_source_tier(url) == 1 for url in ranked_sources) else "普通公开来源待复核"
    )

    if str(sanitized.get("独立生存状态核验", "")) == "确认已故":
        survival_evidence = str(sanitized.get("生存状态核验依据", ""))
        survival_urls = _extract_urls(survival_evidence)
        authoritative_domains = {
            _source_domain(url) for url in survival_urls if _is_official_profile_url(url)
        }
        institution_match = _has_identity_context(source_row, survival_evidence)
        has_death_date = bool(
            re.search(
                r"(?:19|20)\d{2}|\b(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\b"
                r"|\d{1,2}月\d{1,2}日|享年\s*\d+|aged?\s+\d+",
                survival_evidence,
                flags=re.IGNORECASE,
            )
        )
        death_confirmed = (
            sanitized.get("补全身份锁定") == "已锁定"
            and institution_match
            and len(authoritative_domains) >= 2
            and has_death_date
        )
        if not death_confirmed:
            # 共享核验层可能命中同名讣告。补全分支没有同时满足身份、机构、
            # 日期和双权威来源时，不向交付表传播“疑似已故”这一高伤害结论。
            sanitized["独立生存状态核验"] = "待核验"
            sanitized["生存状态"] = ""
            sanitized["生存状态核验依据"] = (
                "共享核验线索未通过补全分支身份闸门，不作为死亡结论："
                "确认死亡需身份锁定、机构匹配、明确死亡日期，"
                "并至少具有两个独立的大学/政府权威来源。"
                + survival_evidence
            )

    return sanitized


def _evidence_payload(row: dict[str, Any]) -> dict[str, str]:
    payload = {}
    for field in EVIDENCE_FIELDS + ["补全身份锁定", "补全身份锁定依据", "补全来源等级"]:
        if _is_placeholder(row.get(field)):
            continue
        limit = 2400 if field == "补全定向证据" else 900
        payload[field] = _clip(row.get(field), limit)
    return payload


def _source_rows_by_number(rows: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    indexed = {}
    for row in rows:
        try:
            indexed[int(row["__excel_row__"])] = row
        except (KeyError, TypeError, ValueError):
            continue
    return indexed


def _direct_updates(
    source_row: dict[str, Any],
    evidence_row: dict[str, Any],
    columns: list[str],
) -> dict[str, str]:
    updates: dict[str, str] = {}
    source_hints: dict[str, list[str]] = {}
    identity_locked = evidence_row.get("补全身份锁定") == "已锁定"
    ranked_sources = _ranked_sources(evidence_row, limit=20)
    official_profile_sources = [
        url for url in ranked_sources if _is_official_profile_url(url)
    ]
    for column in columns:
        if not _is_placeholder(source_row.get(column)):
            continue
        role = _header_role(column)
        if not role:
            continue
        if role == "H指数" and evidence_row.get("补全身份锁定") != "已锁定":
            continue
        for evidence_field in DIRECT_FILL_PRIORITY.get(role, []):
            candidate = evidence_row.get(evidence_field)
            if _is_placeholder(candidate) or _is_low_quality_completion_value(column, candidate):
                continue
            if role in NARRATIVE_ROLES:
                # 叙述字段可直接复用外部增强层已整理出的事实，但必须先锁定身份，
                # 并至少保留一个可供验证员核对的公开来源。
                if not identity_locked:
                    continue
                hints = list(official_profile_sources)
                if role in {"研究兴趣", "领域关联依据", "主要成果", "入选依据"}:
                    hints.extend(
                        url for url in ranked_sources
                        if _source_domain(url) in {"openalex.org", "semanticscholar.org"}
                    )
                hints = list(dict.fromkeys(hints))[:3]
                if not hints:
                    continue
                source_hints[column] = hints
            updates[column] = _clip(candidate, 1800)
            break
    if source_hints:
        evidence_row["补全字段来源提示"] = json.dumps(source_hints, ensure_ascii=False)
    return updates


def _extract_json_array(content: str) -> list[dict[str, Any]]:
    text = str(content or "").strip().replace("```json", "").replace("```", "")
    start = text.find("[")
    end = text.rfind("]")
    if start < 0 or end < start:
        return []
    parsed = json.loads(text[start:end + 1])
    return parsed if isinstance(parsed, list) else []


def _validated_updates(
    raw_items: list[dict[str, Any]],
    source_rows: list[dict[str, Any]],
    columns: list[str],
    update_key: str,
) -> list[dict[str, Any]]:
    source_map = _source_rows_by_number(source_rows)
    allowed_columns = set(columns)
    cleaned = []
    for item in raw_items:
        try:
            row_number = int(item.get("row_number"))
        except (TypeError, ValueError):
            continue
        source_row = source_map.get(row_number)
        if not source_row:
            continue
        raw_updates = item.get(update_key, {})
        if not isinstance(raw_updates, dict):
            continue
        updates = {
            str(column): _clip(value, 2000)
            for column, value in raw_updates.items()
            if column in allowed_columns
            and _is_placeholder(source_row.get(column))
            and not _is_placeholder(value)
            and not _is_low_quality_completion_value(str(column), value)
        }
        if updates:
            cleaned.append({"row_number": row_number, update_key: updates})
    return cleaned


def _chunks(items: list[Any], size: int) -> list[list[Any]]:
    return [items[index:index + size] for index in range(0, len(items), size)]


def _run_completion_llm_batches(
    stage_prefix: str,
    prompt: str,
    payload: list[dict[str, Any]],
    batch_size: int,
) -> list[dict[str, Any]]:
    """补全分支专用的缩批、重试和单行降级，不影响主检索 LLM 调用。"""
    attempts = max(1, int(os.environ.get("TABLE_COMPLETION_LLM_RETRY_ATTEMPTS", "2")))
    retry_delay = max(
        0.0,
        float(os.environ.get("TABLE_COMPLETION_LLM_RETRY_DELAY_SECONDS", "2")),
    )
    results: list[dict[str, Any]] = []

    def invoke_batch(batch: list[dict[str, Any]], label: str) -> bool:
        last_error: Exception | None = None
        for attempt in range(1, attempts + 1):
            try:
                response = invoke_llm_with_stage(
                    label,
                    [
                        SystemMessage(content=prompt),
                        HumanMessage(content=json.dumps(batch, ensure_ascii=False, default=str)),
                    ],
                )
                results.extend(_extract_json_array(response.content))
                return True
            except Exception as exc:
                last_error = exc
                safe_print(
                    f"[{stage_prefix}] {label} 第 {attempt}/{attempts} 次失败：{exc}"
                )
                if attempt < attempts and retry_delay:
                    time.sleep(retry_delay * attempt)
        if len(batch) > 1:
            safe_print(f"[{stage_prefix}] {label} 自动降级为逐位专家验证。")
            for row_index, row in enumerate(batch, start=1):
                invoke_batch([row], f"{label}-单人{row_index}")
            return True
        safe_print(
            f"[{stage_prefix}] {label} 单行调用仍失败，跳过该行并继续整张表：{last_error}"
        )
        return False

    for index, batch in enumerate(_chunks(payload, max(1, batch_size)), start=1):
        invoke_batch(batch, f"{stage_prefix}第{index}批")
    return results


def _gate_validation_results(
    raw_results: list[dict[str, Any]],
    source_rows: list[dict[str, Any]],
    columns: list[str],
    evidence_rows: list[dict[str, Any]] | None = None,
) -> list[dict[str, Any]]:
    source_map = _source_rows_by_number(source_rows)
    evidence_map = _source_rows_by_number(evidence_rows or [])
    allowed_columns = set(columns)
    gated: list[dict[str, Any]] = []
    for item in raw_results:
        try:
            row_number = int(item.get("row_number"))
        except (TypeError, ValueError):
            continue
        source_row = source_map.get(row_number)
        if source_row is None:
            continue
        accepted = item.get("accepted_updates", {})
        field_sources = item.get("field_sources", {})
        rejected = dict(item.get("rejected_columns", {}) or {})
        if not isinstance(accepted, dict):
            accepted = {}
        if not isinstance(field_sources, dict):
            field_sources = {}

        approved: dict[str, str] = {}
        approved_sources: dict[str, list[str]] = {}
        evidence_row = evidence_map.get(row_number, {})
        known_urls = set(_ranked_sources(evidence_row, limit=30))
        for column, value in accepted.items():
            if (
                column not in allowed_columns
                or not _is_placeholder(source_row.get(column))
                or _is_low_quality_completion_value(column, value)
            ):
                continue
            urls = field_sources.get(column, [])
            if isinstance(urls, str):
                urls = _extract_urls(urls)
            urls = [url for url in urls if isinstance(url, str) and _source_tier(url) < 4]
            urls = list(dict.fromkeys(urls))[:3]
            urls = [url for url in urls if url in known_urls]
            if not urls:
                rejected[column] = "缺少本行证据中可核验的中高可信公开来源 URL"
                continue
            role = _header_role(column)
            identity_status = evidence_row.get("补全身份锁定", "待复核")
            normalized_column = _normalize_header(column)
            narrative_like = role in NARRATIVE_ROLES or any(
                token in normalized_column for token in ["类型", "备注"]
            )
            high_risk_field = role in {"H指数", "入选依据", "生存状态"} or "顶尖" in normalized_column
            if high_risk_field and identity_status != "已锁定":
                rejected[column] = "身份尚未锁定，不写入指标、头衔、奖项或死亡状态"
                continue
            if identity_status != "已锁定" and narrative_like:
                homepage_domains = {
                    _source_domain(url)
                    for url in _extract_urls(evidence_row.get("个人主页", ""))
                    if _source_domain(url)
                }
                source_domains = {_source_domain(url) for url in urls if _source_domain(url)}
                has_authoritative_source = any(_source_tier(url) <= 2 for url in urls)
                has_verified_homepage_source = (
                    str(evidence_row.get("成功访问主页", "")).strip() == "是"
                    and bool(source_domains & homepage_domains)
                )
                if not has_authoritative_source and not has_verified_homepage_source:
                    rejected[column] = "身份待复核时，叙述字段缺少权威来源或已验证个人主页支持"
                    continue
            if any(token in normalized_column for token in ["性别", "是否华人"]):
                if not all(_source_tier(url) <= 2 for url in urls):
                    rejected[column] = "敏感身份字段仅接受权威机构或组织来源的明确陈述"
                    continue
                if not _explicit_sensitive_value_supported(column, value, evidence_row, urls):
                    rejected[column] = "来源未明确陈述该敏感属性，禁止根据姓名、照片或代词推断"
                    continue
            if (role in {"入选依据", "生存状态"} or "顶尖" in normalized_column) and not any(
                _source_tier(url) <= 2 for url in urls
            ):
                rejected[column] = "头衔、奖项或生存状态缺少权威来源"
                continue
            approved[column] = _clip(value, 2000)
            approved_sources[column] = urls
        gated.append(
            {
                "row_number": row_number,
                "accepted_updates": approved,
                "field_sources": approved_sources,
                "rejected_columns": rejected,
            }
        )
    return gated


def _add_derived_completion_fields(
    validation_results: list[dict[str, Any]],
    source_rows: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
    columns: list[str],
) -> list[dict[str, Any]]:
    """只基于已通过字段生成可审计的推荐理由，并同步重复的顶尖标记。"""
    source_map = _source_rows_by_number(source_rows)
    evidence_map = _source_rows_by_number(evidence_rows)
    recommendation_columns = [column for column in columns if _header_role(column) == "入选依据"]
    top_columns = [column for column in columns if "顶尖" in _normalize_header(column)]

    for result in validation_results:
        try:
            row_number = int(result.get("row_number"))
        except (TypeError, ValueError):
            continue
        source_row = source_map.get(row_number, {})
        evidence_row = evidence_map.get(row_number, {})
        accepted = result.setdefault("accepted_updates", {})
        field_sources = result.setdefault("field_sources", {})
        if evidence_row.get("补全身份锁定") != "已锁定":
            continue

        accepted_by_role: dict[str, tuple[str, str]] = {}
        for column, value in accepted.items():
            role = _header_role(column)
            if role:
                accepted_by_role[role] = (column, str(value))

        for column in recommendation_columns:
            if not _is_placeholder(source_row.get(column)) or column in accepted:
                continue
            position = accepted_by_role.get("职位")
            interests = accepted_by_role.get("研究兴趣")
            if not position and not interests:
                continue
            supporting_columns = [item[0] for item in [position, interests] if item]
            sources = list(
                dict.fromkeys(
                    url
                    for supporting_column in supporting_columns
                    for url in field_sources.get(supporting_column, [])
                    if _source_tier(url) <= 2
                )
            )[:3]
            if not sources:
                continue
            if position and interests:
                reason = (
                    f"现任{position[1]}，研究聚焦{interests[1]}，"
                    "具有明确的领域相关性与专业代表性。"
                )
            elif position:
                reason = f"现任{position[1]}，具有可核验的专业任职与领域代表性。"
            else:
                reason = f"长期研究{interests[1]}，与目标领域具有明确关联。"
            accepted[column] = _clip(reason, 420)
            field_sources[column] = sources

        accepted_top = next(
            (
                column for column in top_columns
                if str(accepted.get(column, "")).strip() == "是"
                and field_sources.get(column)
            ),
            None,
        )
        if not accepted_top:
            high_signal_sources: list[str] = []
            for candidate_column, candidate_value in accepted.items():
                if not re.search(
                    r"院士|国家科学院|国家医学院|诺贝尔|图灵奖|Gruber|"
                    r"National Academy|Nobel|Turing|\bFellow\b|\bPrize\b|\bMedal\b",
                    str(candidate_value),
                    flags=re.IGNORECASE,
                ):
                    continue
                high_signal_sources.extend(
                    url for url in field_sources.get(candidate_column, [])
                    if _source_tier(url) <= 2
                )
            high_signal_sources = list(dict.fromkeys(high_signal_sources))[:3]
            if high_signal_sources:
                for column in top_columns:
                    if _is_placeholder(source_row.get(column)):
                        accepted[column] = "是"
                        field_sources[column] = high_signal_sources
                accepted_top = top_columns[0] if top_columns else None
        if accepted_top:
            for column in top_columns:
                if column == accepted_top or not _is_placeholder(source_row.get(column)):
                    continue
                accepted[column] = "是"
                field_sources[column] = list(field_sources[accepted_top])[:3]
    return validation_results


def _run_dataframe_stage(
    label: str,
    dataframe: pd.DataFrame,
    operation: Callable[[pd.DataFrame], pd.DataFrame],
) -> pd.DataFrame:
    try:
        result = operation(dataframe)
        safe_print(f"[表格补全工具] {label} 完成，处理 {len(result)} 行。")
        return result
    except Exception as exc:
        safe_print(f"[表格补全工具警告] {label} 失败，保留已有证据并继续：{exc}")
        return dataframe


def _attach_completion_openalex_institution(
    dataframe: pd.DataFrame,
    query: str,
) -> pd.DataFrame:
    """仅为表格补全身份锁提供 OpenAlex 机构，不改变共享增强层。"""
    rows = []
    for _, row in dataframe.iterrows():
        row_data = row.to_dict()
        name = str(row_data.get("专家姓名", "") or "").strip()
        existing_id = str(row_data.get("OpenAlex作者ID", "") or "").strip()
        if name and existing_id:
            try:
                metrics = get_openalex_author_metrics(name, query)
            except Exception as exc:
                safe_print(f"[表格补全身份锁警告] {name} OpenAlex机构读取失败：{exc}")
                metrics = {}
            if str(metrics.get("openalex_id", "") or "").strip() == existing_id:
                row_data["OpenAlex匹配机构"] = metrics.get("institution", "")
        rows.append(row_data)
    return pd.DataFrame(rows)


def _run_external_enrichment(canonical_rows: list[dict[str, Any]], query: str) -> list[dict[str, Any]]:
    dataframe = pd.DataFrame(canonical_rows)
    dataframe = _run_dataframe_stage(
        "OpenAlex",
        dataframe,
        lambda current: enrich_openalex_metrics(current, query),
    )
    dataframe = _run_dataframe_stage(
        "OpenAlex机构身份锁",
        dataframe,
        lambda current: _attach_completion_openalex_institution(current, query),
    )
    enable_s2 = os.environ.get("TABLE_COMPLETION_ENABLE_SEMANTIC_SCHOLAR", "false").lower() == "true"
    if enable_s2:
        dataframe = _run_dataframe_stage(
            "Semantic Scholar",
            dataframe,
            lambda current: enrich_semantic_scholar_metrics(current, query),
        )
    else:
        safe_print("[表格补全工具] Semantic Scholar 默认按需关闭，避免 1 req/s 限流拖慢整批补全。")

    old_opencli_limit = os.environ.get("OPENCLI_HOMEPAGE_MAX_CALLS_PER_PROCESS")
    old_s2_retry_limit = os.environ.get("FINAL_ENRICHMENT_MAX_S2_ROWS")
    completion_overrides = {
        "EXPERT_ENRICHMENT_EXHAUSTIVE_MAX_HOMEPAGE_RETRY_ROWS": os.environ.get(
            "EXPERT_ENRICHMENT_EXHAUSTIVE_MAX_HOMEPAGE_RETRY_ROWS"
        ),
        "EXPERT_ENRICHMENT_EXHAUSTIVE_MAX_TAVILY_ROWS": os.environ.get(
            "EXPERT_ENRICHMENT_EXHAUSTIVE_MAX_TAVILY_ROWS"
        ),
    }
    try:
        if not enable_s2:
            os.environ["FINAL_ENRICHMENT_MAX_S2_ROWS"] = "0"
        # 表格补全稍后会按专家执行一次综合定向检索。这里关闭共享增强层中
        # 重复的多查询 Tavily 回退，避免长表在前几轮耗尽搜索额度。
        os.environ["EXPERT_ENRICHMENT_EXHAUSTIVE_MAX_HOMEPAGE_RETRY_ROWS"] = os.environ.get(
            "TABLE_COMPLETION_HOMEPAGE_RETRY_ROWS_PER_ROUND", "0"
        )
        os.environ["EXPERT_ENRICHMENT_EXHAUSTIVE_MAX_TAVILY_ROWS"] = os.environ.get(
            "TABLE_COMPLETION_LEGACY_TAVILY_ROWS_PER_ROUND", "0"
        )
        dataframe = _run_dataframe_stage(
            "主页、OpenCLI 与 Tavily 定向补全",
            dataframe,
            lambda current: enrich_expert_details(current, query, exhaustive=True),
        )
    finally:
        if old_opencli_limit is None:
            os.environ.pop("OPENCLI_HOMEPAGE_MAX_CALLS_PER_PROCESS", None)
        else:
            os.environ["OPENCLI_HOMEPAGE_MAX_CALLS_PER_PROCESS"] = old_opencli_limit
        if old_s2_retry_limit is None:
            os.environ.pop("FINAL_ENRICHMENT_MAX_S2_ROWS", None)
        else:
            os.environ["FINAL_ENRICHMENT_MAX_S2_ROWS"] = old_s2_retry_limit
        for name, old_value in completion_overrides.items():
            if old_value is None:
                os.environ.pop(name, None)
            else:
                os.environ[name] = old_value

    dataframe = _run_dataframe_stage("联系方式清洗", dataframe, clean_contact_fields)
    dataframe = _run_dataframe_stage(
        "主页访问状态规范化",
        dataframe,
        normalize_homepage_access_status,
    )
    dataframe = _run_dataframe_stage(
        "独立生存状态核验",
        dataframe,
        lambda current: verify_survival_status(current, query, exhaustive=True),
    )
    dataframe = dataframe.drop(columns=["专家姓名验证", "姓名验证依据"], errors="ignore")
    dataframe = add_name_verification_columns(dataframe)
    return dataframe.fillna("").to_dict(orient="records")


COMPLETION_RESEARCHER_PROMPT = """
你是“已有专家表格信息补全研究员”。Python 已用 OpenAlex、Semantic Scholar、公开网页、
专家主页、OpenCLI 和 Tavily 为每一行收集证据。你的任务不是发现或新增专家，而是只根据
输入 evidence 为原表中的空白字段提出候选值。

规则：
1. 只能填写 missing_columns 中的字段，禁止修改任何已有值，禁止新增或删除专家。
2. 不得根据姓名、照片或语言猜测性别、国籍、族裔、“是否华人”等敏感属性；仅在证据明确写明时填写。
3. 邮箱、主页、单位、职位、履历、成果、推荐理由必须有输入证据支持；无证据则不输出该字段。
4. 不把检索摘要、错误信息或“暂无公开信息”写入表格。
5. 叙述字段统一输出简洁、完整的中文；不得直接复制截断的英文搜索摘要。
6. 补全身份锁定为“未通过”时，不得采用数据库指标、履历、奖项或生存状态信息；
   为“待复核”时不得填写 H 指数、引用数、头衔、奖项或死亡状态等高风险字段。
7. 优先利用“补全定向证据”填写教育经历、职务、工作经历、细分领域、推荐理由、人才类型和备注。
8. “主要教育经历”至少应包含学校或培训机构，不能只写“博士学位”；证据只有学位而无学校时宁可不填。
9. “推荐理由”应概括有来源支持的职位、代表成果、奖项或领域贡献；“类型”可依据已验证职位归类为
   高校教授、科研机构研究员、临床专家、企业技术领军者等。备注只做已验证事实摘要。
10. “是否推荐为顶尖/是否为顶尖人才”只有在院士、权威大奖、Fellow、重要学术领导职位或明确高影响力
    指标得到权威来源支持时填写“是”；证据不足时不填，不得猜测。
11. 只输出 JSON 数组，格式：
   [{"row_number": 5, "updates": {"职务": "教授", "主要工作经历": "..."}}]
"""

COMPLETION_GAP_FILL_PROMPT = COMPLETION_RESEARCHER_PROMPT + """

这是第二轮缺失字段定向整理。第一轮后仍为空的字段已经列在 missing_columns 中。
请逐项检查补全定向证据和已验证的职位、教育、工作经历、研究方向之间的关系，优先补齐：
职务、完整教育经历、完整工作经历、细分领域、推荐理由、人才类型和备注。
可以把同一来源明确支持的多项事实整理成中文，但仍不得猜测敏感属性或虚构荣誉。
"""

COMPLETION_VALIDATOR_PROMPT = """
你是“已有专家表格补全验证员”。请逐项核对 proposed_updates 是否被 evidence 直接支持，
并检查是否发生同名错配、机构错配、无来源推断或敏感属性推断。

规则：
1. 只有证据充分且身份一致的值才能进入 accepted_updates。
2. 无法确认的字段必须拒绝，不得为了完整率降低真实性标准。
3. 每个通过字段必须在 field_sources 中列出 1 至 3 个直接支持该字段的 URL；没有 URL 的叙述字段不得通过。
4. 院士、Fellow、重要奖项和死亡状态必须由大学、政府、学会、奖项官网等权威来源直接支持。
5. evidence 中的“补全字段来源提示”是 Python 根据身份锁定和来源等级生成的候选 URL；
   仍需核对证据内容，但应优先使用这些 URL，避免漏掉外部工具已经整理出的职位、履历和研究方向。
6. proposed_updates 中的每个字段都必须明确归入 accepted_updates 或 rejected_columns，不得静默遗漏。
7. 只输出 JSON 数组：
   [{"row_number": 5, "accepted_updates": {"职务": "教授"},
     "field_sources": {"职务": ["https://example.edu/profile"]},
     "rejected_columns": {"性别": "证据未明确"}}]
"""

COMPLETION_CORRECTOR_PROMPT = """
你是“已有专家表格补全纠错员”。请结合原始行、证据、研究员建议和验证反馈，输出最终可回填值。

规则：
1. 默认采用 accepted_updates；被拒绝字段只有在 evidence 中存在明确直接证据时才可修正后恢复。
2. 禁止覆盖原有非空单元格，禁止新增或删除专家，禁止猜测。
3. 所有叙述字段须改写为自然、完整、简洁的中文，不得保留截断英文或搜索摘要腔。
4. 只输出 JSON 数组：
   [{"row_number": 5, "final_updates": {"职务": "教授"}}]
"""


def completion_researcher_node(state: TableCompletionState) -> dict[str, Any]:
    source_map = _source_rows_by_number(state["source_rows"])
    raw_evidence_rows = _run_external_enrichment(state["canonical_rows"], state["query"])
    completion_tavily = _make_completion_tavily()
    evidence_rows = []
    for evidence_row in raw_evidence_rows:
        try:
            row_number = int(evidence_row["__excel_row__"])
        except (KeyError, TypeError, ValueError):
            continue
        source_row = source_map.get(row_number)
        if source_row is not None:
            evidence_row = _add_targeted_completion_evidence(
                evidence_row,
                source_row,
                state["query"],
                completion_tavily,
            )
            evidence_rows.append(_sanitize_completion_evidence(evidence_row, source_row))
    batch_size = max(1, int(os.environ.get("TABLE_COMPLETION_LLM_BATCH_SIZE", "10")))
    payload_rows = []
    direct_by_row: dict[int, dict[str, str]] = {}

    for evidence_row in evidence_rows:
        try:
            row_number = int(evidence_row["__excel_row__"])
        except (KeyError, TypeError, ValueError):
            continue
        source_row = source_map.get(row_number)
        if source_row is None:
            continue
        direct = _direct_updates(source_row, evidence_row, state["columns"])
        direct_by_row[row_number] = direct
        missing_columns = [
            column
            for column in state["columns"]
            if _is_placeholder(source_row.get(column)) and column not in direct
        ]
        if missing_columns:
            payload_rows.append(
                {
                    "row_number": row_number,
                    "expert": source_row.get("__source_name__", ""),
                    "existing": {
                        column: _clip(source_row.get(column), 500)
                        for column in state["columns"]
                        if not _is_placeholder(source_row.get(column))
                    },
                    "missing_columns": missing_columns,
                    "evidence": _evidence_payload(evidence_row),
                }
            )

    researcher_batch_size = max(
        1,
        int(os.environ.get("TABLE_COMPLETION_RESEARCHER_BATCH_SIZE", str(min(batch_size, 5)))),
    )
    llm_updates = _run_completion_llm_batches(
        "表格补全研究员",
        COMPLETION_RESEARCHER_PROMPT,
        payload_rows,
        researcher_batch_size,
    )

    cleaned_llm = _validated_updates(
        llm_updates,
        state["source_rows"],
        state["columns"],
        "updates",
    )
    llm_field_count = sum(len(item["updates"]) for item in cleaned_llm)
    direct_field_count = sum(len(updates) for updates in direct_by_row.values())
    merged_by_row = {item["row_number"]: dict(item["updates"]) for item in cleaned_llm}
    for row_number, direct in direct_by_row.items():
        merged_by_row.setdefault(row_number, {}).update(direct)

    enable_second_pass = os.environ.get(
        "TABLE_COMPLETION_ENABLE_SECOND_PASS", "true"
    ).lower() == "true"
    if enable_second_pass:
        evidence_map = _source_rows_by_number(evidence_rows)
        gap_payload = []
        max_rows = max(0, int(os.environ.get("TABLE_COMPLETION_SECOND_PASS_MAX_ROWS", "40")))
        for row_number, source_row in source_map.items():
            remaining = [
                column
                for column in state["columns"]
                if _is_placeholder(source_row.get(column))
                and column not in merged_by_row.get(row_number, {})
            ]
            if not remaining:
                continue
            gap_payload.append(
                {
                    "row_number": row_number,
                    "expert": source_row.get("__source_name__", ""),
                    "existing": {
                        **{
                            column: _clip(source_row.get(column), 500)
                            for column in state["columns"]
                            if not _is_placeholder(source_row.get(column))
                        },
                        **merged_by_row.get(row_number, {}),
                    },
                    "missing_columns": remaining,
                    "evidence": _evidence_payload(evidence_map.get(row_number, {})),
                }
            )
        gap_batch_size = max(
            1,
            int(os.environ.get("TABLE_COMPLETION_SECOND_PASS_BATCH_SIZE", "3")),
        )
        gap_updates = _run_completion_llm_batches(
            "表格补全研究员缺失字段",
            COMPLETION_GAP_FILL_PROMPT,
            gap_payload[:max_rows],
            gap_batch_size,
        )
        cleaned_gaps = _validated_updates(
            gap_updates,
            state["source_rows"],
            state["columns"],
            "updates",
        )
        gap_field_count = 0
        for item in cleaned_gaps:
            target = merged_by_row.setdefault(item["row_number"], {})
            for column, value in item["updates"].items():
                if column not in target:
                    target[column] = value
                    gap_field_count += 1
        safe_print(f"[表格补全研究员] 第二轮新增 {gap_field_count} 个待验证字段建议。")
    research_updates = [
        {"row_number": row_number, "updates": updates}
        for row_number, updates in merged_by_row.items()
        if updates
    ]
    safe_print(
        f"[表格补全研究员] 外部工具处理 {len(evidence_rows)} 行，"
        f"为 {len(research_updates)} 行提出可验证补全建议；"
        f"确定性候选 {direct_field_count} 项，LLM 首轮候选 {llm_field_count} 项。"
    )
    return {"evidence_rows": evidence_rows, "research_updates": research_updates}


def completion_validator_node(state: TableCompletionState) -> dict[str, Any]:
    proposals = state.get("research_updates", [])
    if not proposals:
        return {"validation_results": []}
    source_map = _source_rows_by_number(state["source_rows"])
    evidence_map = _source_rows_by_number(state.get("evidence_rows", []))
    batch_size = max(
        1,
        int(os.environ.get("TABLE_COMPLETION_VALIDATOR_BATCH_SIZE", "2")),
    )
    payload = [
        {
            "row_number": item["row_number"],
            "original": source_map[item["row_number"]],
            "proposed_updates": item["updates"],
            "evidence": _evidence_payload(evidence_map.get(item["row_number"], {})),
        }
        for item in proposals
    ]
    validation_results = _run_completion_llm_batches(
        "表格补全验证员",
        COMPLETION_VALIDATOR_PROMPT,
        payload,
        batch_size,
    )
    gated_results = _gate_validation_results(
        validation_results,
        state["source_rows"],
        state["columns"],
        state.get("evidence_rows", []),
    )
    gated_results = _add_derived_completion_fields(
        gated_results,
        state["source_rows"],
        state.get("evidence_rows", []),
        state["columns"],
    )
    accepted_field_count = sum(
        len(item.get("accepted_updates", {})) for item in gated_results
    )
    accepted_by_column: dict[str, int] = {}
    for item in gated_results:
        for column in item.get("accepted_updates", {}):
            accepted_by_column[column] = accepted_by_column.get(column, 0) + 1
    accepted_summary = "、".join(
        f"{column}={count}" for column, count in sorted(accepted_by_column.items())
    ) or "无"
    safe_print(
        f"[表格补全验证员] 字段级证据闸门批准 {accepted_field_count} 个字段；"
        f"字段分布：{accepted_summary}。"
    )
    return {"validation_results": gated_results}


def completion_corrector_node(state: TableCompletionState) -> dict[str, Any]:
    validations = state.get("validation_results", [])
    if not validations:
        return {"final_updates": []}
    source_map = _source_rows_by_number(state["source_rows"])
    evidence_map = _source_rows_by_number(state.get("evidence_rows", []))
    proposal_map = {
        int(item["row_number"]): item["updates"]
        for item in state.get("research_updates", [])
    }
    batch_size = max(
        1,
        int(os.environ.get("TABLE_COMPLETION_CORRECTOR_BATCH_SIZE", "3")),
    )
    payload = []
    for validation in validations:
        try:
            row_number = int(validation.get("row_number"))
        except (TypeError, ValueError):
            continue
        if row_number not in source_map:
            continue
        payload.append(
            {
                "row_number": row_number,
                "original": source_map[row_number],
                "evidence": _evidence_payload(evidence_map.get(row_number, {})),
                "researcher_updates": proposal_map.get(row_number, {}),
                "validation": validation,
            }
        )

    raw_final = _run_completion_llm_batches(
        "表格补全纠错员",
        COMPLETION_CORRECTOR_PROMPT,
        payload,
        batch_size,
    )
    corrected_updates = _validated_updates(
        raw_final,
        state["source_rows"],
        state["columns"],
        "final_updates",
    )
    corrected_by_row = {
        int(item["row_number"]): item["final_updates"]
        for item in corrected_updates
    }
    accepted_updates_by_row = {
        int(item.get("row_number")): dict(item.get("accepted_updates", {}))
        for item in validations
        if str(item.get("row_number", "")).isdigit()
        and isinstance(item.get("accepted_updates", {}), dict)
    }
    gated_updates = []
    for row_number, accepted in accepted_updates_by_row.items():
        # 纠错员负责润色或修正已通过字段；即使它漏回某项，验证员已经批准的
        # 值也必须保留，避免一次非结构化输出造成信息丢失。
        approved = dict(accepted)
        for column, value in corrected_by_row.get(row_number, {}).items():
            if column in accepted and not _is_low_quality_completion_value(column, value):
                approved[column] = value
        if approved:
            gated_updates.append(
                {"row_number": row_number, "final_updates": approved}
            )
    final_updates = gated_updates
    safe_print(f"[表格补全纠错员] 最终批准回填 {len(final_updates)} 行。")
    return {"final_updates": final_updates}


def build_table_completion_graph():
    workflow = StateGraph(TableCompletionState)
    workflow.add_node("completion_researcher", completion_researcher_node)
    workflow.add_node("completion_validator", completion_validator_node)
    workflow.add_node("completion_corrector", completion_corrector_node)
    workflow.set_entry_point("completion_researcher")
    workflow.add_edge("completion_researcher", "completion_validator")
    workflow.add_edge("completion_validator", "completion_corrector")
    workflow.add_edge("completion_corrector", END)
    return workflow.compile()


table_completion_graph = build_table_completion_graph()


def _row_sources(evidence_row: dict[str, Any]) -> str:
    return " | ".join(_ranked_sources(evidence_row, limit=5))


def write_completed_workbook(
    workbook_bytes: bytes,
    layout: WorkbookLayout,
    final_updates: list[dict[str, Any]],
    evidence_rows: list[dict[str, Any]],
    output_path: str,
    validation_results: list[dict[str, Any]] | None = None,
) -> dict[str, int]:
    workbook = load_workbook(io.BytesIO(workbook_bytes), data_only=False)
    worksheet = workbook[layout.sheet_name]
    column_index = {header: index for index, header in layout.columns}
    source_map = _source_rows_by_number(layout.data_rows)
    evidence_map = _source_rows_by_number(evidence_rows)
    last_source_column = max(index for index, _ in layout.columns)
    audit_start = last_source_column + 1

    for offset, header in enumerate(AUDIT_COLUMNS):
        cell = worksheet.cell(layout.header_row, audit_start + offset, header)
        template = worksheet.cell(layout.header_row, last_source_column)
        if template.has_style:
            cell._style = copy(template._style)
        cell.font = copy(template.font)
        cell.fill = copy(template.fill)
        cell.border = copy(template.border)
        cell.alignment = copy(template.alignment)
        worksheet.column_dimensions[get_column_letter(audit_start + offset)].width = 28

    filled_cells = 0
    completed_rows = 0
    updates_map = {
        int(item["row_number"]): item.get("final_updates", {})
        for item in final_updates
    }
    validation_map = {
        int(item["row_number"]): item
        for item in (validation_results or [])
        if str(item.get("row_number", "")).isdigit()
    }
    detail_rows: list[list[Any]] = []
    for row_number, source_row in source_map.items():
        row_filled = 0
        validation = validation_map.get(row_number, {})
        field_sources = validation.get("field_sources", {})
        for column, value in updates_map.get(row_number, {}).items():
            target_column = column_index.get(column)
            if not target_column:
                continue
            cell = worksheet.cell(row_number, target_column)
            if not _is_placeholder(cell.value) or _is_placeholder(value):
                continue
            cell.value = value
            row_filled += 1
            filled_cells += 1
            urls = field_sources.get(column, []) if isinstance(field_sources, dict) else []
            if isinstance(urls, str):
                urls = _extract_urls(urls)
            detail_rows.append(
                [
                    row_number,
                    source_row.get("__source_name__", ""),
                    column,
                    value,
                    "通过字段级证据闸门",
                    " | ".join(urls[:3]),
                    evidence_map.get(row_number, {}).get("补全身份锁定", "待复核"),
                ]
            )
        if row_filled:
            completed_rows += 1

        evidence = evidence_map.get(row_number, {})
        accepted_count = len(validation.get("accepted_updates", {}))
        identity_status = str(evidence.get("补全身份锁定", "待复核") or "待复核")
        audit_values = [
            _row_sources(evidence) or "未找到可写入的中高可信公开来源 URL",
            f"身份{identity_status}；字段级验证通过 {accepted_count} 项",
            (
                f"本次补全 {row_filled} 个空白字段；"
                f"生存状态：{evidence.get('独立生存状态核验', '待核验')}；"
                "原有非空单元格未被覆盖。"
            ),
        ]
        for offset, value in enumerate(audit_values):
            target = worksheet.cell(row_number, audit_start + offset, value)
            template = worksheet.cell(row_number, last_source_column)
            if template.has_style:
                target._style = copy(template._style)
            target.alignment = copy(template.alignment)

    detail_title = "补全证据明细"
    if detail_title in workbook.sheetnames:
        del workbook[detail_title]
    detail_sheet = workbook.create_sheet(detail_title)
    detail_headers = ["原表行号", "专家姓名", "补全字段", "补全值", "验证结果", "字段来源", "身份锁定"]
    detail_sheet.append(detail_headers)
    for row in detail_rows:
        detail_sheet.append(row)
    header_template = worksheet.cell(layout.header_row, last_source_column)
    for cell in detail_sheet[1]:
        cell.font = copy(header_template.font)
        cell.fill = copy(header_template.fill)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, width in enumerate([12, 24, 20, 48, 22, 54, 16], start=1):
        detail_sheet.column_dimensions[get_column_letter(index)].width = width
    detail_sheet.freeze_panes = "A2"
    detail_sheet.auto_filter.ref = detail_sheet.dimensions
    for row in detail_sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return {"total_rows": len(layout.data_rows), "completed_rows": completed_rows, "filled_cells": filled_cells}


def _safe_filename(value: str) -> str:
    cleaned = re.sub(r'[<>:"/\\|?*\x00-\x1f]+', "_", str(value or "")).strip(" ._")
    return cleaned or "专家信息表"


def _run_completion_chunks(
    query: str,
    columns: list[str],
    source_rows: list[dict[str, Any]],
    canonical_rows: list[dict[str, Any]],
    progress_callback: Callable[[str], None] | None = None,
) -> dict[str, list[dict[str, Any]]]:
    """每轮处理固定人数，逐轮运行独立补全图并验证全表覆盖。"""
    chunk_size = max(
        1,
        int(os.environ.get("TABLE_COMPLETION_PROCESS_CHUNK_SIZE", "15")),
    )
    attempts = max(
        1,
        int(os.environ.get("TABLE_COMPLETION_PROCESS_CHUNK_RETRY_ATTEMPTS", "2")),
    )
    canonical_map = _source_rows_by_number(canonical_rows)
    source_chunks = _chunks(source_rows, chunk_size)
    aggregated = {
        "evidence_rows": [],
        "research_updates": [],
        "validation_results": [],
        "final_updates": [],
    }
    processed_row_numbers: set[int] = set()

    for chunk_index, source_chunk in enumerate(source_chunks, start=1):
        row_numbers = [int(row["__excel_row__"]) for row in source_chunk]
        canonical_chunk = [
            canonical_map[row_number]
            for row_number in row_numbers
            if row_number in canonical_map
        ]
        if len(canonical_chunk) != len(source_chunk):
            missing = sorted(set(row_numbers) - {int(row["__excel_row__"]) for row in canonical_chunk})
            raise RuntimeError(f"第 {chunk_index} 轮缺少规范化专家行：{missing}")
        if progress_callback:
            progress_callback(
                f"正在进行第 {chunk_index}/{len(source_chunks)} 轮专家信息补全，"
                f"本轮 {len(source_chunk)} 位专家..."
            )
        last_error: Exception | None = None
        final_state: dict[str, Any] | None = None
        for attempt in range(1, attempts + 1):
            try:
                final_state = table_completion_graph.invoke(
                    {
                        "query": query,
                        "columns": columns,
                        "source_rows": source_chunk,
                        "canonical_rows": canonical_chunk,
                    }
                )
                evidence_numbers = {
                    int(row["__excel_row__"])
                    for row in final_state.get("evidence_rows", [])
                    if row.get("__excel_row__") is not None
                }
                if evidence_numbers != set(row_numbers):
                    missing = sorted(set(row_numbers) - evidence_numbers)
                    extra = sorted(evidence_numbers - set(row_numbers))
                    raise RuntimeError(
                        f"本轮证据覆盖不完整，缺少={missing}，异常行={extra}"
                    )
                break
            except Exception as exc:
                last_error = exc
                safe_print(
                    f"[表格补全轮次] 第 {chunk_index}/{len(source_chunks)} 轮 "
                    f"第 {attempt}/{attempts} 次失败：{exc}"
                )
                if attempt < attempts:
                    time.sleep(min(6, attempt * 2))
        if final_state is None:
            raise RuntimeError(
                f"第 {chunk_index}/{len(source_chunks)} 轮补全连续失败：{last_error}"
            )
        for key in aggregated:
            aggregated[key].extend(final_state.get(key, []))
        processed_row_numbers.update(row_numbers)
        round_updates = final_state.get("final_updates", [])
        updated_rows = {
            int(item["row_number"])
            for item in round_updates
            if item.get("row_number") is not None and item.get("final_updates")
        }
        filled_fields = sum(
            len(item.get("final_updates", {}))
            for item in round_updates
            if isinstance(item.get("final_updates"), dict)
        )
        safe_print(
            f"[表格补全轮次] 第 {chunk_index}/{len(source_chunks)} 轮完成，"
            f"本轮有更新专家={len(updated_rows)}/{len(source_chunk)}，"
            f"通过字段={filled_fields}；累计处理 "
            f"{min(chunk_index * chunk_size, len(source_rows))}/{len(source_rows)} 人。"
        )
    expected_row_numbers = {int(row["__excel_row__"]) for row in source_rows}
    if processed_row_numbers != expected_row_numbers:
        missing = sorted(expected_row_numbers - processed_row_numbers)
        raise RuntimeError(f"补全轮次结束后仍有专家未处理：{missing}")
    return aggregated


def run_table_completion(
    workbook_bytes: bytes,
    original_filename: str,
    domain_hint: str = "",
    output_dir: str = "Expert_Results",
    progress_callback: Callable[[str], None] | None = None,
) -> tuple[str | None, str | None, dict[str, int] | None]:
    """运行独立的已有专家表格补全分支。"""
    try:
        if progress_callback:
            progress_callback("正在识别 Excel 表头和专家数据行...")
        layout = inspect_workbook_layout(workbook_bytes)
        configure_tavily_budget(len(layout.data_rows))
        canonical_rows = build_canonical_rows(layout)
        query = domain_hint.strip() or Path(original_filename).stem
        if progress_callback:
            round_size = max(
                1,
                int(os.environ.get("TABLE_COMPLETION_PROCESS_CHUNK_SIZE", "15")),
            )
            total_rounds = (len(layout.data_rows) + round_size - 1) // round_size
            progress_callback(
                f"已识别 {len(layout.data_rows)} 位专家，将按每轮 {round_size} 位执行 "
                f"{total_rounds} 轮补全，正在调用外部工具和三个补全智能体..."
            )
        final_state = _run_completion_chunks(
            query,
            [header for _, header in layout.columns],
            layout.data_rows,
            canonical_rows,
            progress_callback,
        )
        timestamp = time.strftime("%Y%m%d_%H%M%S")
        stem = _safe_filename(Path(original_filename).stem)
        output_path = str(Path(output_dir).resolve() / f"{stem}_{timestamp}_信息补全.xlsx")
        summary = write_completed_workbook(
            workbook_bytes,
            layout,
            final_state.get("final_updates", []),
            final_state.get("evidence_rows", []),
            output_path,
            final_state.get("validation_results", []),
        )
        if progress_callback:
            progress_callback("补全结果已完成验证并写入 Excel。")
        return output_path, None, summary
    except Exception as exc:
        safe_print(f"[表格信息补全失败] {exc}")
        return None, f"表格信息补全失败：{exc}", None
